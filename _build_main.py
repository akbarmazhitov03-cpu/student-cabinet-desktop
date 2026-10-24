from pathlib import Path

root = Path(r"C:\Users\user\Desktop\student_cabinet")

main_code = '''from __future__ import annotations

import argparse
import sqlite3
from datetime import datetime
from pathlib import Path
from tkinter import StringVar, Tk, filedialog, messagebox, ttk
from typing import Any

from argon2 import PasswordHasher
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

DB_PATH = Path(__file__).resolve().with_name("student_cabinet.db")
APP_TITLE = "Student Cabinet"

TRANSLATIONS = {
    "ru": {
        "login": "Вход",
        "username": "Логин",
        "password": "Пароль",
        "signin": "Войти",
        "logout": "Выйти",
        "dashboard": "Главная",
        "profile": "Профиль",
        "grades": "Оценки",
        "schedule": "Расписание",
        "attendance": "Посещаемость",
        "payments": "Платежи",
        "messages": "Сообщения",
        "students": "Студенты",
        "search": "Поиск",
        "add": "Добавить",
        "save": "Сохранить",
        "error": "Ошибка",
        "success": "Успешно",
        "welcome": "Добро пожаловать",
        "no_data": "Нет данных",
        "change_password": "Сменить пароль",
        "current_password": "Текущий пароль",
        "new_password": "Новый пароль",
        "invalid_login": "Неверный логин или пароль.",
        "admin_hint": "Администратор: admin / admin123456",
        "student_hint": "Студент: student1 / student123",
        "language": "Язык",
        "mark_read": "Отметить прочитанным",
        "export": "Экспорт",
        "password_changed": "Пароль изменён",
    },
    "tg": {
        "login": "Воридшавӣ",
        "username": "Логин",
        "password": "Рамз",
        "signin": "Даромадан",
        "logout": "Баромадан",
        "dashboard": "Асосӣ",
        "profile": "Профил",
        "grades": "Баҳоҳо",
        "schedule": "Ҷадвал",
        "attendance": "Иштирок",
        "payments": "Пардохт",
        "messages": "Паёмҳо",
        "students": "Донишҷӯён",
        "search": "Ҷустуҷӯ",
        "add": "Илова",
        "save": "Захира",
        "error": "Хато",
        "success": "Муваффақият",
        "welcome": "Хуш омадед",
        "no_data": "Маълумоти нест",
        "change_password": "Иваз кардани рамз",
        "current_password": "Рамзи ҷорӣ",
        "new_password": "Рамзи нав",
        "invalid_login": "Логин ё рамзи нодуруст.",
        "admin_hint": "Маъмур: admin / admin123456",
        "student_hint": "Донишҷӯ: student1 / student123",
        "language": "Забон",
        "mark_read": "Хондани баррасӣ",
        "export": "Содирот",
        "password_changed": "Рамз иваз шуд",
    },
}


def t(key: str, lang: str = "ru") -> str:
    return TRANSLATIONS.get(lang, TRANSLATIONS["ru"]).get(key, key)


class AuthenticationService:
    """Password hashing and verification using Argon2."""

    _hasher = PasswordHasher()

    @staticmethod
    def hash_password(password: str) -> str:
        if len(password) < 8:
            raise ValueError("Password must be at least 8 characters long.")
        return AuthenticationService._hasher.hash(password)

    @staticmethod
    def verify_password(password_hash: str, password: str) -> bool:
        try:
            return AuthenticationService._hasher.verify(password_hash, password)
        except Exception:
            return False


class DatabaseManager:
    """SQLite-backed local database manager for the desktop-only Student Cabinet."""

    def __init__(self, db_path: str | Path | None = None) -> None:
        self.db_path = Path(db_path) if db_path else DB_PATH
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self.conn = sqlite3.connect(str(self.db_path))
        self.conn.row_factory = sqlite3.Row
        self.conn.execute("PRAGMA foreign_keys = ON")
        self.init_schema()
        self.ensure_seed_data()

    def init_schema(self) -> None:
        self.conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT NOT NULL UNIQUE,
                password_hash TEXT NOT NULL,
                role TEXT NOT NULL CHECK(role IN ('admin', 'student')),
                is_active INTEGER NOT NULL DEFAULT 1,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );

            CREATE TABLE IF NOT EXISTS groups (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE
            );

            CREATE TABLE IF NOT EXISTS specialties (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE
            );

            CREATE TABLE IF NOT EXISTS courses (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE,
                year_number INTEGER NOT NULL UNIQUE
            );

            CREATE TABLE IF NOT EXISTS semesters (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                academic_year TEXT NOT NULL,
                term INTEGER NOT NULL,
                starts_on TEXT NOT NULL,
                ends_on TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS students (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL UNIQUE,
                student_number TEXT NOT NULL UNIQUE,
                full_name TEXT NOT NULL,
                group_id INTEGER NOT NULL,
                specialty_id INTEGER,
                course_id INTEGER,
                email TEXT,
                phone TEXT,
                birth_date TEXT,
                FOREIGN KEY(user_id) REFERENCES users(id),
                FOREIGN KEY(group_id) REFERENCES groups(id),
                FOREIGN KEY(specialty_id) REFERENCES specialties(id),
                FOREIGN KEY(course_id) REFERENCES courses(id)
            );

            CREATE TABLE IF NOT EXISTS subjects (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE,
                code TEXT NOT NULL UNIQUE
            );

            CREATE TABLE IF NOT EXISTS enrollments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL,
                subject_id INTEGER NOT NULL,
                semester_id INTEGER NOT NULL,
                UNIQUE(student_id, subject_id, semester_id),
                FOREIGN KEY(student_id) REFERENCES students(id),
                FOREIGN KEY(subject_id) REFERENCES subjects(id),
                FOREIGN KEY(semester_id) REFERENCES semesters(id)
            );

            CREATE TABLE IF NOT EXISTS grades (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL,
                subject_id INTEGER NOT NULL,
                semester_id INTEGER NOT NULL,
                current_score REAL,
                examination_score REAL,
                final_score REAL,
                note TEXT,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(student_id, subject_id, semester_id),
                FOREIGN KEY(student_id) REFERENCES students(id),
                FOREIGN KEY(subject_id) REFERENCES subjects(id),
                FOREIGN KEY(semester_id) REFERENCES semesters(id)
            );

            CREATE TABLE IF NOT EXISTS schedule (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                group_id INTEGER NOT NULL,
                subject_id INTEGER NOT NULL,
                weekday INTEGER NOT NULL CHECK(weekday BETWEEN 0 AND 6),
                starts_at TEXT NOT NULL,
                ends_at TEXT NOT NULL,
                teacher TEXT NOT NULL,
                classroom TEXT NOT NULL,
                semester_id INTEGER,
                FOREIGN KEY(group_id) REFERENCES groups(id),
                FOREIGN KEY(subject_id) REFERENCES subjects(id),
                FOREIGN KEY(semester_id) REFERENCES semesters(id)
            );

            CREATE TABLE IF NOT EXISTS attendance (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL,
                subject_id INTEGER NOT NULL,
                lesson_date TEXT NOT NULL,
                status TEXT NOT NULL CHECK(status IN ('present', 'absent', 'excused')),
                notes TEXT,
                UNIQUE(student_id, subject_id, lesson_date),
                FOREIGN KEY(student_id) REFERENCES students(id),
                FOREIGN KEY(subject_id) REFERENCES subjects(id)
            );

            CREATE TABLE IF NOT EXISTS contracts (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL,
                amount REAL NOT NULL CHECK(amount >= 0),
                contract_number TEXT NOT NULL UNIQUE,
                academic_year TEXT NOT NULL,
                FOREIGN KEY(student_id) REFERENCES students(id)
            );

            CREATE TABLE IF NOT EXISTS payments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                contract_id INTEGER NOT NULL,
                amount REAL NOT NULL CHECK(amount > 0),
                payment_date TEXT NOT NULL,
                receipt_number TEXT NOT NULL UNIQUE,
                note TEXT,
                FOREIGN KEY(contract_id) REFERENCES contracts(id)
            );

            CREATE TABLE IF NOT EXISTS messages (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                author_id INTEGER NOT NULL,
                target_group_id INTEGER,
                title TEXT NOT NULL,
                body TEXT NOT NULL,
                published_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                is_published INTEGER NOT NULL DEFAULT 1,
                FOREIGN KEY(author_id) REFERENCES users(id),
                FOREIGN KEY(target_group_id) REFERENCES groups(id)
            );

            CREATE TABLE IF NOT EXISTS message_reads (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                message_id INTEGER NOT NULL,
                user_id INTEGER NOT NULL,
                read_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(message_id, user_id),
                FOREIGN KEY(message_id) REFERENCES messages(id),
                FOREIGN KEY(user_id) REFERENCES users(id)
            );

            CREATE TABLE IF NOT EXISTS audit_logs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                actor_user_id INTEGER,
                action TEXT NOT NULL,
                entity_type TEXT NOT NULL,
                entity_id TEXT,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY(actor_user_id) REFERENCES users(id)
            );
            """
        )
        self.conn.commit()

    def log_audit(self, actor_user_id: int | None, action: str, entity_type: str, entity_id: str | None = None) -> None:
        self.conn.execute(
            "INSERT INTO audit_logs (actor_user_id, action, entity_type, entity_id, created_at) VALUES (?, ?, ?, ?, ?)",
            (actor_user_id, action, entity_type, entity_id, datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        )
        self.conn.commit()

    def ensure_seed_data(self) -> None:
        if self.conn.execute("SELECT COUNT(*) FROM users").fetchone()[0] > 0:
            return

        self.conn.execute(
            "INSERT INTO users (username, password_hash, role, is_active) VALUES (?, ?, 'admin', 1)",
            ("admin", AuthenticationService.hash_password("admin123456")),
        )
        admin_id = self.conn.execute("SELECT id FROM users WHERE username = 'admin'").fetchone()[0]

        self.conn.execute("INSERT INTO groups (name) VALUES ('G-101')")
        group_id = self.conn.execute("SELECT id FROM groups WHERE name = 'G-101'").fetchone()[0]
        self.conn.execute("INSERT INTO specialties (name) VALUES ('Computer Science')")
        specialty_id = self.conn.execute("SELECT id FROM specialties WHERE name = 'Computer Science'").fetchone()[0]
        self.conn.execute("INSERT INTO courses (name, year_number) VALUES ('FIRST', 1)")
        course_id = self.conn.execute("SELECT id FROM courses WHERE name = 'FIRST'").fetchone()[0]
        self.conn.execute(
            "INSERT INTO semesters (name, academic_year, term, starts_on, ends_on) VALUES ('Autumn', '2025/2026', 1, '2025-09-01', '2025-12-31')"
        )
        semester_id = self.conn.execute("SELECT id FROM semesters WHERE name = 'Autumn'").fetchone()[0]

        self.conn.execute(
            "INSERT INTO users (username, password_hash, role, is_active) VALUES (?, ?, 'student', 1)",
            ("student1", AuthenticationService.hash_password("student123")),
        )
        student_user_id = self.conn.execute("SELECT id FROM users WHERE username = 'student1'").fetchone()[0]

        self.conn.execute(
            "INSERT INTO students (user_id, student_number, full_name, group_id, specialty_id, course_id, email, phone, birth_date) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (student_user_id, "ST-1001", "Ali Rahmonov", group_id, specialty_id, course_id, "ali@example.com", "+992900000001", "2006-03-15"),
        )
        student_id = self.conn.execute("SELECT id FROM students WHERE student_number = 'ST-1001'").fetchone()[0]

        self.conn.execute("INSERT INTO subjects (name, code) VALUES ('Mathematics', 'MATH101')")
        subject_id = self.conn.execute("SELECT id FROM subjects WHERE code = 'MATH101'").fetchone()[0]
        self.conn.execute(
            "INSERT INTO grades (student_id, subject_id, semester_id, current_score, examination_score, final_score, note) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (student_id, subject_id, semester_id, 88.0, 92.0, 90.0, "Good"),
        )
        self.conn.execute(
            "INSERT INTO attendance (student_id, subject_id, lesson_date, status, notes) VALUES (?, ?, ?, ?, ?)",
            (student_id, subject_id, "2025-10-10", "present", "Active"),
        )
        self.conn.execute(
            "INSERT INTO schedule (group_id, subject_id, weekday, starts_at, ends_at, teacher, classroom, semester_id) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            (group_id, subject_id, 1, "09:00", "10:30", "Mr. Saidov", "A101", semester_id),
        )
        self.conn.execute(
            "INSERT INTO contracts (student_id, amount, contract_number, academic_year) VALUES (?, ?, ?, ?)",
            (student_id, 2500.0, "CT-1001", "2025/2026"),
        )
        contract_id = self.conn.execute("SELECT id FROM contracts WHERE contract_number = 'CT-1001'").fetchone()[0]
        self.conn.execute(
            "INSERT INTO payments (contract_id, amount, payment_date, receipt_number, note) VALUES (?, ?, ?, ?, ?)",
            (contract_id, 1200.0, "2025-09-15", "RCPT-1001", "Initial payment"),
        )
        self.conn.execute(
            "INSERT INTO messages (author_id, target_group_id, title, body, published_at, is_published) VALUES (?, ?, ?, ?, ?, 1)",
            (admin_id, group_id, "Welcome", "Welcome to Student Cabinet!", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        )
        self.conn.commit()


class StudentService:
    @staticmethod
    def get_by_user_id(conn: sqlite3.Connection, user_id: int) -> dict[str, Any] | None:
        row = conn.execute(
            """
            SELECT s.*, u.username, g.name AS group_name, sp.name AS specialty_name
            FROM students s
            JOIN users u ON u.id = s.user_id
            LEFT JOIN groups g ON g.id = s.group_id
            LEFT JOIN specialties sp ON sp.id = s.specialty_id
            WHERE u.id = ?
            """,
            (user_id,),
        ).fetchone()
        return dict(row) if row else None

    @staticmethod
    def list(conn: sqlite3.Connection, query: str = "") -> list[dict[str, Any]]:
        sql = """
            SELECT s.*, u.username, g.name AS group_name, sp.name AS specialty_name
            FROM students s
            JOIN users u ON u.id = s.user_id
            LEFT JOIN groups g ON g.id = s.group_id
            LEFT JOIN specialties sp ON sp.id = s.specialty_id
        """
        params: list[Any] = []
        if query:
            term = f"%{query.strip()}%"
            sql += " WHERE s.full_name LIKE ? OR s.student_number LIKE ? OR u.username LIKE ?"
            params.extend([term, term, term])
        sql += " ORDER BY s.full_name"
        return [dict(row) for row in conn.execute(sql, params).fetchall()]

    @staticmethod
    def create(conn: sqlite3.Connection, payload: dict[str, Any]) -> int:
        username = (payload.get("username") or payload.get("student_number") or "student").strip()
        password = (payload.get("password") or "student123").strip()
        if not username:
            raise ValueError("Username is required.")
        if len(password) < 8:
            raise ValueError("Password must be at least 8 characters long.")
        user_id = conn.execute(
            "INSERT INTO users (username, password_hash, role, is_active) VALUES (?, ?, 'student', 1)",
            (username, AuthenticationService.hash_password(password)),
        ).lastrowid
        student_id = conn.execute(
            "INSERT INTO students (user_id, student_number, full_name, group_id, specialty_id, course_id, email, phone, birth_date) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (
                user_id,
                payload.get("student_number") or f"ST-{user_id:04d}",
                payload.get("full_name") or username,
                int(payload.get("group_id") or 1),
                int(payload["specialty_id"]) if payload.get("specialty_id") else None,
                int(payload["course_id"]) if payload.get("course_id") else None,
                payload.get("email") or None,
                payload.get("phone") or None,
                payload.get("birth_date") or None,
            ),
        ).lastrowid
        conn.commit()
        return int(student_id)


class GradeService:
    @staticmethod
    def list_for_student(conn: sqlite3.Connection, student_id: int) -> list[dict[str, Any]]:
        rows = conn.execute(
            """
            SELECT g.*, s.name AS subject_name, sem.name AS semester_name
            FROM grades g
            JOIN subjects s ON s.id = g.subject_id
            JOIN semesters sem ON sem.id = g.semester_id
            WHERE g.student_id = ?
            ORDER BY sem.name
            """,
            (student_id,),
        ).fetchall()
        return [dict(row) for row in rows]

    @staticmethod
    def list_all(conn: sqlite3.Connection, query: str = "") -> list[dict[str, Any]]:
        sql = """
            SELECT g.*, s.name AS subject_name, sem.name AS semester_name,
                   st.full_name, st.student_number
            FROM grades g
            JOIN subjects s ON s.id = g.subject_id
            JOIN semesters sem ON sem.id = g.semester_id
            JOIN students st ON st.id = g.student_id
        """
        params: list[Any] = []
        if query:
            term = f"%{query.strip()}%"
            sql += " WHERE st.full_name LIKE ? OR s.name LIKE ? OR st.student_number LIKE ?"
            params.extend([term, term, term])
        sql += " ORDER BY st.full_name, s.name"
        return [dict(row) for row in conn.execute(sql, params).fetchall()]

    @staticmethod
    def save(conn: sqlite3.Connection, payload: dict[str, Any]) -> None:
        if not payload.get("student_id") or not payload.get("subject_id") or not payload.get("semester_id"):
            raise ValueError("Student, subject and semester are required.")
        current_score = float(payload.get("current_score") or 0)
        exam_score = float(payload.get("examination_score") or 0)
        final_value = payload.get("final_score")
        if final_value in (None, ""):
            final_score = round((current_score + exam_score) / 2, 2)
        else:
            final_score = float(final_value)
        if not 0 <= current_score <= 100 or not 0 <= exam_score <= 100 or not 0 <= final_score <= 100:
            raise ValueError("Scores must be between 0 and 100.")
        conn.execute(
            """
            INSERT INTO grades (student_id, subject_id, semester_id, current_score, examination_score, final_score, note, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(student_id, subject_id, semester_id)
            DO UPDATE SET
                current_score = excluded.current_score,
                examination_score = excluded.examination_score,
                final_score = excluded.final_score,
                note = excluded.note,
                updated_at = CURRENT_TIMESTAMP
            """,
            (
                int(payload["student_id"]),
                int(payload["subject_id"]),
                int(payload["semester_id"]),
                current_score,
                exam_score,
                final_score,
                payload.get("note") or None,
            ),
        )
        conn.commit()


class ScheduleService:
    @staticmethod
    def list_for_student(conn: sqlite3.Connection, group_id: int) -> list[dict[str, Any]]:
        rows = conn.execute(
            """
            SELECT sc.*, subj.name AS subject_name
            FROM schedule sc
            JOIN subjects subj ON subj.id = sc.subject_id
            WHERE sc.group_id = ?
            ORDER BY sc.weekday, sc.starts_at
            """,
            (group_id,),
        ).fetchall()
        return [dict(row) for row in rows]

    @staticmethod
    def list_all(conn: sqlite3.Connection, query: str = "") -> list[dict[str, Any]]:
        sql = """
            SELECT sc.*, g.name AS group_name, subj.name AS subject_name
            FROM schedule sc
            JOIN groups g ON g.id = sc.group_id
            JOIN subjects subj ON subj.id = sc.subject_id
        """
        params: list[Any] = []
        if query:
            term = f"%{query.strip()}%"
            sql += " WHERE g.name LIKE ? OR subj.name LIKE ? OR sc.teacher LIKE ?"
            params.extend([term, term, term])
        sql += " ORDER BY g.name, sc.weekday, sc.starts_at"
        return [dict(row) for row in conn.execute(sql, params).fetchall()]


class AttendanceService:
    @staticmethod
    def list_for_student(conn: sqlite3.Connection, student_id: int) -> list[dict[str, Any]]:
        rows = conn.execute(
            """
            SELECT a.*, s.name AS subject_name
            FROM attendance a
            JOIN subjects s ON s.id = a.subject_id
            WHERE a.student_id = ?
            ORDER BY a.lesson_date DESC
            """,
            (student_id,),
        ).fetchall()
        return [dict(row) for row in rows]

    @staticmethod
    def list_all(conn: sqlite3.Connection, query: str = "") -> list[dict[str, Any]]:
        sql = """
            SELECT a.*, s.name AS subject_name, st.full_name, st.student_number
            FROM attendance a
            JOIN subjects s ON s.id = a.subject_id
            JOIN students st ON st.id = a.student_id
        """
        params: list[Any] = []
        if query:
            term = f"%{query.strip()}%"
            sql += " WHERE st.full_name LIKE ? OR s.name LIKE ? OR st.student_number LIKE ?"
            params.extend([term, term, term])
        sql += " ORDER BY a.lesson_date DESC, st.full_name"
        return [dict(row) for row in conn.execute(sql, params).fetchall()]

    @staticmethod
    def attendance_percentage(rows: list[dict[str, Any]]) -> float | None:
        if not rows:
            return None
        valid = [row for row in rows if row.get("status") != "excused"]
        if not valid:
            return None
        present = sum(1 for row in valid if row.get("status") == "present")
        return round((present / len(valid)) * 100, 2)


class PaymentService:
    @staticmethod
    def list_for_student(conn: sqlite3.Connection, student_id: int) -> list[dict[str, Any]]:
        rows = conn.execute(
            """
            SELECT c.id, c.contract_number, c.academic_year, c.amount,
                   COALESCE(SUM(p.amount), 0) AS paid
            FROM contracts c
            LEFT JOIN payments p ON p.contract_id = c.id
            WHERE c.student_id = ?
            GROUP BY c.id, c.contract_number, c.academic_year, c.amount
            ORDER BY c.academic_year DESC
            """,
            (student_id,),
        ).fetchall()
        return [
            {
                "contract_number": row["contract_number"],
                "academic_year": row["academic_year"],
                "amount": float(row["amount"]),
                "paid": round(float(row["paid"]), 2),
                "outstanding": round(float(row["amount"]) - float(row["paid"]), 2),
            }
            for row in rows
        ]

    @staticmethod
    def list_all(conn: sqlite3.Connection, query: str = "") -> list[dict[str, Any]]:
        sql = """
            SELECT p.*, c.contract_number, st.full_name
            FROM payments p
            JOIN contracts c ON c.id = p.contract_id
            JOIN students st ON st.id = c.student_id
        """
        params: list[Any] = []
        if query:
            term = f"%{query.strip()}%"
            sql += " WHERE st.full_name LIKE ? OR c.contract_number LIKE ? OR p.receipt_number LIKE ?"
            params.extend([term, term, term])
        sql += " ORDER BY p.payment_date DESC"
        return [dict(row) for row in conn.execute(sql, params).fetchall()]

    @staticmethod
    def create_payment(conn: sqlite3.Connection, payload: dict[str, Any]) -> None:
        contract_id = int(payload["contract_id"])
        amount = float(payload.get("amount") or 0)
        if amount <= 0:
            raise ValueError("Payment amount must be positive.")
        contract = conn.execute("SELECT amount FROM contracts WHERE id = ?", (contract_id,)).fetchone()
        if contract is None:
            raise ValueError("Contract not found.")
        total_paid = float(conn.execute("SELECT COALESCE(SUM(amount), 0) FROM payments WHERE contract_id = ?", (contract_id,)).fetchone()[0])
        outstanding = float(contract["amount"]) - total_paid
        if amount > outstanding:
            raise ValueError("Payment exceeds outstanding balance.")
        conn.execute(
            "INSERT INTO payments (contract_id, amount, payment_date, receipt_number, note) VALUES (?, ?, ?, ?, ?)",
            (
                contract_id,
                amount,
                payload.get("payment_date") or datetime.now().strftime("%Y-%m-%d"),
                payload.get("receipt_number") or f"RCP-{datetime.now().strftime('%Y%m%d%H%M%S')}",
                payload.get("note") or None,
            ),
        )
        conn.commit()


class MessageService:
    @staticmethod
    def list_for_student(conn: sqlite3.Connection, user_id: int, student: dict[str, Any]) -> list[dict[str, Any]]:
        rows = conn.execute(
            """
            SELECT m.*, mr.id AS read_id
            FROM messages m
            LEFT JOIN message_reads mr ON mr.message_id = m.id AND mr.user_id = ?
            WHERE m.is_published = 1 AND (m.target_group_id IS NULL OR m.target_group_id = ?)
            ORDER BY m.published_at DESC
            """,
            (user_id, student.get("group_id")),
        ).fetchall()
        return [
            {
                "id": row["id"],
                "title": row["title"],
                "body": row["body"],
                "published_at": row["published_at"],
                "read": row["read_id"] is not None,
            }
            for row in rows
        ]

    @staticmethod
    def list_all(conn: sqlite3.Connection, query: str = "") -> list[dict[str, Any]]:
        sql = """
            SELECT m.*, g.name AS group_name
            FROM messages m
            LEFT JOIN groups g ON g.id = m.target_group_id
        """
        params: list[Any] = []
        if query:
            term = f"%{query.strip()}%"
            sql += " WHERE m.title LIKE ? OR m.body LIKE ? OR g.name LIKE ?"
            params.extend([term, term, term])
        sql += " ORDER BY m.published_at DESC"
        return [dict(row) for row in conn.execute(sql, params).fetchall()]

    @staticmethod
    def mark_read(conn: sqlite3.Connection, message_id: int, user_id: int) -> None:
        if conn.execute("SELECT 1 FROM message_reads WHERE message_id = ? AND user_id = ?", (message_id, user_id)).fetchone() is not None:
            return
        conn.execute(
            "INSERT INTO message_reads (message_id, user_id, read_at) VALUES (?, ?, ?)",
            (message_id, user_id, datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        )
        conn.commit()


class ExcelExporter:
    @staticmethod
    def export_rows(rows: list[dict[str, Any]], output_path: str | Path, sheet_name: str = "Report") -> Path:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name
        if rows:
            headers = list(rows[0].keys())
            sheet.append(headers)
            for row in rows:
                sheet.append([row.get(header) for header in headers])
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(fill_type="solid", fgColor="17365D")
        else:
            sheet.append(["No data"])
        target = Path(output_path)
        target.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(target)
        return target


class APIClient:
    """Placeholder client for future network integrations; local mode does not use it."""

    def __init__(self, base_url: str = "http://localhost") -> None:
        self.base_url = base_url

    def get(self, path: str) -> dict[str, Any]:
        raise RuntimeError("This desktop build does not use a network API.")

    def post(self, path: str, payload: dict[str, Any]) -> dict[str, Any]:
        raise RuntimeError("This desktop build does not use a network API.")


class LoginWindow(ttk.Frame):
    def __init__(self, root: Tk, app: "StudentCabinetApp") -> None:
        super().__init__(root)
        self.app = app
        self.pack(fill="both", expand=True)

        card = ttk.Frame(self, padding=30)
        card.place(relx=0.5, rely=0.5, anchor="center")

        ttk.Label(card, text=APP_TITLE, font=("Segoe UI", 22, "bold")).grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 16))
        ttk.Label(card, text=t("language", self.app.lang)).grid(row=1, column=0, sticky="w")
        self.lang_var = StringVar(value=self.app.lang)
        combo = ttk.Combobox(card, textvariable=self.lang_var, values=["ru", "tg"], state="readonly", width=10)
        combo.grid(row=1, column=1, sticky="e", pady=(0, 8))
        combo.bind("<<ComboboxSelected>>", lambda event: self.app.set_language(self.lang_var.get()))

        ttk.Label(card, text=t("username", self.app.lang)).grid(row=2, column=0, sticky="w")
        self.username_var = StringVar(value="admin")
        ttk.Entry(card, textvariable=self.username_var).grid(row=3, column=0, columnspan=2, sticky="ew", pady=(0, 8))

        ttk.Label(card, text=t("password", self.app.lang)).grid(row=4, column=0, sticky="w")
        self.password_var = StringVar(value="admin123456")
        ttk.Entry(card, textvariable=self.password_var, show="*").grid(row=5, column=0, columnspan=2, sticky="ew", pady=(0, 12))

        ttk.Button(card, text=t("signin", self.app.lang), command=self.submit).grid(row=6, column=0, columnspan=2, sticky="ew")
        self.status = ttk.Label(card, text="", foreground="#B42318")
        self.status.grid(row=7, column=0, columnspan=2, sticky="w", pady=(10, 0))

        ttk.Label(card, text=t("admin_hint", self.app.lang), justify="left").grid(row=8, column=0, columnspan=2, sticky="w", pady=(12, 0))
        ttk.Label(card, text=t("student_hint", self.app.lang), justify="left").grid(row=9, column=0, columnspan=2, sticky="w")

    def submit(self) -> None:
        username = self.username_var.get().strip()
        password = self.password_var.get()
        if not username or not password:
            self.status.configure(text=t("error", self.app.lang))
            return
        try:
            self.app.login(username, password)
        except ValueError as exc:
            self.status.configure(text=str(exc))


class BaseDashboard(ttk.Frame):
    def __init__(self, root: Tk, app: "StudentCabinetApp", user: dict[str, Any]) -> None:
        super().__init__(root)
        self.app = app
        self.user = user
        self.pack(fill="both", expand=True)
        self.columnconfigure(1, weight=1)
        self.rowconfigure(0, weight=1)

        self.sidebar = ttk.Frame(self, width=220)
        self.sidebar.grid(row=0, column=0, sticky="ns")
        self.content = ttk.Frame(self)
        self.content.grid(row=0, column=1, sticky="nsew", padx=12, pady=12)
        self.content.columnconfigure(0, weight=1)

        self.build_sidebar()
        self.show_view("dashboard")

    def build_sidebar(self) -> None:
        ttk.Label(self.sidebar, text="Student\nCabinet", font=("Segoe UI", 18, "bold")).pack(anchor="w", padx=12, pady=(18, 16))
        for name, label in self.menu_items():
            ttk.Button(self.sidebar, text=label, command=lambda n=name: self.show_view(n)).pack(fill="x", padx=10, pady=4)
        ttk.Button(self.sidebar, text=t("change_password", self.app.lang), command=self.change_password).pack(fill="x", padx=10, pady=(18, 4))
        ttk.Button(self.sidebar, text=t("logout", self.app.lang), command=self.app.logout).pack(fill="x", padx=10, pady=4)

    def menu_items(self) -> list[tuple[str, str]]:
        return [
            ("dashboard", t("dashboard", self.app.lang)),
            ("profile", t("profile", self.app.lang)),
            ("grades", t("grades", self.app.lang)),
            ("schedule", t("schedule", self.app.lang)),
            ("attendance", t("attendance", self.app.lang)),
            ("payments", t("payments", self.app.lang)),
            ("messages", t("messages", self.app.lang)),
        ]

    def render_table(self, rows: list[dict[str, Any]], fields: list[str]) -> None:
        for child in self.content.winfo_children():
            child.destroy()
        if not rows:
            ttk.Label(self.content, text=t("no_data", self.app.lang)).pack(anchor="w")
            return
        tree = ttk.Treeview(self.content, columns=fields, show="headings")
        for field in fields:
            tree.heading(field, text=field.replace("_", " ").title())
            tree.column(field, width=160, anchor="w")
        for row in rows:
            values = [str(row.get(field)) if row.get(field) is not None else "" for field in fields]
            tree.insert("", "end", values=values)
        y_scroll = ttk.Scrollbar(self.content, orient="vertical", command=tree.yview)
        x_scroll = ttk.Scrollbar(self.content, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=y_scroll.set, xscrollcommand=x_scroll.set)
        tree.pack(side="left", fill="both", expand=True)
        y_scroll.pack(side="right", fill="y")
        x_scroll.pack(side="bottom", fill="x")

    def show_view(self, view: str) -> None:
        for child in self.content.winfo_children():
            child.destroy()
        if view == "dashboard":
            self.render_dashboard(); return
        student = StudentService.get_by_user_id(self.app.conn, self.user["id"])
        if student is None:
            ttk.Label(self.content, text="Student profile missing.").pack(anchor="w")
            return
        if view == "profile":
            self.render_table([student], ["student_number", "full_name", "email", "phone", "group_name"])
        elif view == "grades":
            self.render_table(GradeService.list_for_student(self.app.conn, student["id"]), ["subject_name", "semester_name", "current_score", "examination_score", "final_score", "note"])
        elif view == "schedule":
            self.render_table(ScheduleService.list_for_student(self.app.conn, student["group_id"]), ["weekday", "starts_at", "ends_at", "subject_name", "teacher", "classroom"])
        elif view == "attendance":
            rows = AttendanceService.list_for_student(self.app.conn, student["id"])
            percentage = AttendanceService.attendance_percentage(rows)
            if percentage is not None:
                ttk.Label(self.content, text=f"Attendance: {percentage}%").pack(anchor="w", pady=(0, 8))
            self.render_table(rows, ["lesson_date", "subject_name", "status", "notes"])
        elif view == "payments":
            self.render_table(PaymentService.list_for_student(self.app.conn, student["id"]), ["contract_number", "academic_year", "amount", "paid", "outstanding"])
        elif view == "messages":
            rows = MessageService.list_for_student(self.app.conn, self.user["id"], student)
            self.render_table(rows, ["title", "published_at", "read", "body"])
            for row in rows:
                if not row["read"]:
                    ttk.Button(self.content, text=t("mark_read", self.app.lang), command=lambda mid=row["id"]: self.mark_read(mid)).pack(anchor="e", pady=8)

    def render_dashboard(self) -> None:
        student = StudentService.get_by_user_id(self.app.conn, self.user["id"])
        name = student["full_name"] if student else self.user["username"]
        ttk.Label(self.content, text=f"{t('welcome', self.app.lang)}, {name}", font=("Segoe UI", 18, "bold")).pack(anchor="w", pady=(0, 18))
        cards = ttk.Frame(self.content)
        cards.pack(fill="x")
        for item in ["grades", "schedule", "attendance", "payments"]:
            frame = ttk.LabelFrame(cards, text=t(item, self.app.lang), padding=12)
            frame.pack(side="left", fill="both", expand=True, padx=(0, 10), pady=4)
            ttk.Button(frame, text=t("dashboard", self.app.lang), command=lambda n=item: self.show_view(n)).pack(anchor="w")

    def mark_read(self, message_id: int) -> None:
        MessageService.mark_read(self.app.conn, message_id, self.user["id"])
        self.show_view("messages")

    def change_password(self) -> None:
        dialog = ttk.Toplevel(self)
        dialog.title(t("change_password", self.app.lang))
        dialog.geometry("330x220")
        ttk.Label(dialog, text=t("current_password", self.app.lang)).pack(anchor="w", padx=12, pady=(12, 0))
        old_var = StringVar()
        ttk.Entry(dialog, textvariable=old_var, show="*").pack(fill="x", padx=12, pady=(0, 8))
        ttk.Label(dialog, text=t("new_password", self.app.lang)).pack(anchor="w", padx=12)
        new_var = StringVar()
        ttk.Entry(dialog, textvariable=new_var, show="*").pack(fill="x", padx=12, pady=(0, 8))

        def save() -> None:
            row = self.app.conn.execute("SELECT password_hash FROM users WHERE id = ?", (self.user["id"],)).fetchone()
            if row is None or not AuthenticationService.verify_password(row["password_hash"], old_var.get()):
                messagebox.showerror(t("error", self.app.lang), "Current password is incorrect.")
                return
            new_password = new_var.get().strip()
            if len(new_password) < 8:
                messagebox.showerror(t("error", self.app.lang), "Password must be at least 8 characters long.")
                return
            self.app.conn.execute("UPDATE users SET password_hash = ? WHERE id = ?", (AuthenticationService.hash_password(new_password), self.user["id"]))
            self.app.conn.commit()
            messagebox.showinfo(t("success", self.app.lang), t("password_changed", self.app.lang))
            dialog.destroy()

        ttk.Button(dialog, text=t("save", self.app.lang), command=save).pack(pady=12)


class StudentDashboard(BaseDashboard):
    pass


class AdminDashboard(BaseDashboard):
    def menu_items(self) -> list[tuple[str, str]]:
        return [
            ("dashboard", t("dashboard", self.app.lang)),
            ("students", t("students", self.app.lang)),
            ("grades", t("grades", self.app.lang)),
            ("schedule", t("schedule", self.app.lang)),
            ("attendance", t("attendance", self.app.lang)),
            ("payments", t("payments", self.app.lang)),
            ("messages", t("messages", self.app.lang)),
        ]

    def show_view(self, view: str) -> None:
        for child in self.content.winfo_children():
            child.destroy()
        if view == "dashboard":
            self.render_dashboard(); return
        toolbar = ttk.Frame(self.content)
        toolbar.pack(fill="x", pady=(0, 8))
        search_var = StringVar()
        ttk.Entry(toolbar, textvariable=search_var).pack(side="left", fill="x", expand=True)
        ttk.Button(toolbar, text=t("search", self.app.lang), command=lambda: self.load_view(view, search_var.get())).pack(side="left", padx=(8, 0))
        ttk.Button(toolbar, text=t("add", self.app.lang), command=lambda: self.add_record(view)).pack(side="left", padx=(8, 0))
        ttk.Button(toolbar, text=t("export", self.app.lang), command=lambda: self.export_report(view)).pack(side="left", padx=(8, 0))
        self.load_view(view, "")

    def load_view(self, view: str, query: str) -> None:
        if view == "students":
            self.render_table(StudentService.list(self.app.conn, query), ["student_number", "full_name", "username", "group_name", "email"])
        elif view == "grades":
            self.render_table(GradeService.list_all(self.app.conn, query), ["full_name", "student_number", "subject_name", "semester_name", "final_score", "note"])
        elif view == "schedule":
            self.render_table(ScheduleService.list_all(self.app.conn, query), ["group_name", "subject_name", "weekday", "starts_at", "ends_at", "teacher", "classroom"])
        elif view == "attendance":
            self.render_table(AttendanceService.list_all(self.app.conn, query), ["full_name", "student_number", "subject_name", "lesson_date", "status", "notes"])
        elif view == "payments":
            self.render_table(PaymentService.list_all(self.app.conn, query), ["full_name", "contract_number", "amount", "payment_date", "receipt_number"])
        elif view == "messages":
            self.render_table(MessageService.list_all(self.app.conn, query), ["title", "group_name", "published_at", "body"])

    def add_record(self, view: str) -> None:
        if view == "students":
            self.add_student_dialog()
        elif view == "grades":
            self.add_grade_dialog()
        elif view == "schedule":
            self.add_schedule_dialog()
        elif view == "attendance":
            self.add_attendance_dialog()
        elif view == "payments":
            self.add_payment_dialog()
        elif view == "messages":
            self.add_message_dialog()

    def build_form(self, parent, fields: list[tuple[str, str]]) -> dict[str, StringVar]:
        values: dict[str, StringVar] = {}
        for index, (key, label) in enumerate(fields):
            ttk.Label(parent, text=label).grid(row=index, column=0, sticky="w", padx=12, pady=6)
            var = StringVar()
            ttk.Entry(parent, textvariable=var).grid(row=index, column=1, sticky="ew", padx=12, pady=6)
            values[key] = var
        return values

    def export_report(self, view: str) -> None:
        mapping = {
            "students": lambda: StudentService.list(self.app.conn),
            "grades": lambda: GradeService.list_all(self.app.conn),
            "schedule": lambda: ScheduleService.list_all(self.app.conn),
            "attendance": lambda: AttendanceService.list_all(self.app.conn),
            "payments": lambda: PaymentService.list_all(self.app.conn),
            "messages": lambda: MessageService.list_all(self.app.conn),
        }
        rows = mapping.get(view, lambda: [])()
        folder = filedialog.askdirectory(title="Export folder")
        if not folder:
            return
        target = Path(folder) / f"{view}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        try:
            ExcelExporter.export_rows(rows, target, sheet_name=view)
            messagebox.showinfo(t("success", self.app.lang), f"Exported: {target}")
        except Exception as exc:
            messagebox.showerror(t("error", self.app.lang), str(exc))

    def add_student_dialog(self) -> None:
        dialog = ttk.Toplevel(self)
        dialog.title("Add student")
        dialog.geometry("480x430")
        fields = [
            ("username", "Username"),
            ("password", "Password"),
            ("student_number", "Student number"),
            ("full_name", "Full name"),
            ("group_id", "Group ID"),
            ("specialty_id", "Specialty ID"),
            ("course_id", "Course ID"),
            ("email", "Email"),
            ("phone", "Phone"),
            ("birth_date", "Birth date YYYY-MM-DD"),
        ]
        values = self.build_form(dialog, fields)

        def save() -> None:
            payload = {key: var.get().strip() for key, var in values.items()}
            if not payload["full_name"] or not payload["student_number"]:
                messagebox.showerror(t("error", self.app.lang), "Student number and full name are required.")
                return
            try:
                StudentService.create(self.app.conn, payload)
                self.app.db.log_audit(self.user["id"], "create_student", "student", payload["student_number"])
            except ValueError as exc:
                messagebox.showerror(t("error", self.app.lang), str(exc))
                return
            messagebox.showinfo(t("success", self.app.lang), "Student created.")
            dialog.destroy(); self.show_view("students")

        ttk.Button(dialog, text=t("save", self.app.lang), command=save).grid(row=len(fields), column=0, columnspan=2, pady=12)

    def add_grade_dialog(self) -> None:
        dialog = ttk.Toplevel(self)
        dialog.title("Add grade")
        dialog.geometry("420x420")
        fields = [
            ("student_id", "Student ID"),
            ("subject_id", "Subject ID"),
            ("semester_id", "Semester ID"),
            ("current_score", "Current score"),
            ("examination_score", "Examination score"),
            ("final_score", "Final score (optional)"),
            ("note", "Note"),
        ]
        values = self.build_form(dialog, fields)

        def save() -> None:
            try:
                GradeService.save(self.app.conn, {
                    "student_id": values["student_id"].get(),
                    "subject_id": values["subject_id"].get(),
                    "semester_id": values["semester_id"].get(),
                    "current_score": values["current_score"].get(),
                    "examination_score": values["examination_score"].get(),
                    "final_score": values["final_score"].get(),
                    "note": values["note"].get(),
                })
                self.app.db.log_audit(self.user["id"], "save_grade", "grade", values["student_id"].get())
            except ValueError as exc:
                messagebox.showerror(t("error", self.app.lang), str(exc))
                return
            dialog.destroy(); self.show_view("grades")

        ttk.Button(dialog, text=t("save", self.app.lang), command=save).grid(row=len(fields), column=0, columnspan=2, pady=12)

    def add_schedule_dialog(self) -> None:
        dialog = ttk.Toplevel(self)
        dialog.title("Add schedule")
        dialog.geometry("450x430")
        fields = [
            ("group_id", "Group ID"),
            ("subject_id", "Subject ID"),
            ("weekday", "Weekday (0-6)"),
            ("starts_at", "Starts HH:MM"),
            ("ends_at", "Ends HH:MM"),
            ("teacher", "Teacher"),
            ("classroom", "Classroom"),
            ("semester_id", "Semester ID"),
        ]
        values = self.build_form(dialog, fields)

        def save() -> None:
            self.app.conn.execute(
                "INSERT INTO schedule (group_id, subject_id, weekday, starts_at, ends_at, teacher, classroom, semester_id) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    int(values["group_id"].get()),
                    int(values["subject_id"].get()),
                    int(values["weekday"].get()),
                    values["starts_at"].get(),
                    values["ends_at"].get(),
                    values["teacher"].get(),
                    values["classroom"].get(),
                    int(values["semester_id"].get()) if values["semester_id"].get() else None,
                ),
            )
            self.app.conn.commit(); self.app.db.log_audit(self.user["id"], "create_schedule", "schedule", values["group_id"].get()); dialog.destroy(); self.show_view("schedule")

        ttk.Button(dialog, text=t("save", self.app.lang), command=save).grid(row=len(fields), column=0, columnspan=2, pady=12)

    def add_attendance_dialog(self) -> None:
        dialog = ttk.Toplevel(self)
        dialog.title("Add attendance")
        dialog.geometry("420x350")
        fields = [
            ("student_id", "Student ID"),
            ("subject_id", "Subject ID"),
            ("lesson_date", "Lesson date YYYY-MM-DD"),
            ("status", "Status"),
            ("notes", "Notes"),
        ]
        values = self.build_form(dialog, fields)

        def save() -> None:
            self.app.conn.execute(
                "INSERT INTO attendance (student_id, subject_id, lesson_date, status, notes) VALUES (?, ?, ?, ?, ?)",
                (
                    int(values["student_id"].get()),
                    int(values["subject_id"].get()),
                    values["lesson_date"].get(),
                    values["status"].get(),
                    values["notes"].get() or None,
                ),
            )
            self.app.conn.commit(); self.app.db.log_audit(self.user["id"], "set_attendance", "attendance", values["student_id"].get()); dialog.destroy(); self.show_view("attendance")

        ttk.Button(dialog, text=t("save", self.app.lang), command=save).grid(row=len(fields), column=0, columnspan=2, pady=12)

    def add_payment_dialog(self) -> None:
        dialog = ttk.Toplevel(self)
        dialog.title("Add payment")
        dialog.geometry("420x350")
        fields = [
            ("contract_id", "Contract ID"),
            ("amount", "Amount"),
            ("payment_date", "Date YYYY-MM-DD"),
            ("receipt_number", "Receipt number"),
            ("note", "Note"),
        ]
        values = self.build_form(dialog, fields)

        def save() -> None:
            try:
                PaymentService.create_payment(self.app.conn, {
                    "contract_id": values["contract_id"].get(),
                    "amount": values["amount"].get(),
                    "payment_date": values["payment_date"].get(),
                    "receipt_number": values["receipt_number"].get(),
                    "note": values["note"].get(),
                })
                self.app.db.log_audit(self.user["id"], "add_payment", "payment", values["contract_id"].get())
            except ValueError as exc:
                messagebox.showerror(t("error", self.app.lang), str(exc))
                return
            dialog.destroy(); self.show_view("payments")

        ttk.Button(dialog, text=t("save", self.app.lang), command=save).grid(row=len(fields), column=0, columnspan=2, pady=12)

    def add_message_dialog(self) -> None:
        dialog = ttk.Toplevel(self)
        dialog.title("Add message")
        dialog.geometry("420x300")
        fields = [
            ("target_group_id", "Target group ID (optional)"),
            ("title", "Title"),
            ("body", "Body"),
        ]
        values = self.build_form(dialog, fields)

        def save() -> None:
            self.app.conn.execute(
                "INSERT INTO messages (author_id, target_group_id, title, body, published_at, is_published) VALUES (?, ?, ?, ?, ?, 1)",
                (
                    self.user["id"],
                    int(values["target_group_id"].get()) if values["target_group_id"].get() else None,
                    values["title"].get(),
                    values["body"].get(),
                    datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                ),
            )
            self.app.conn.commit(); self.app.db.log_audit(self.user["id"], "send_message", "message", values["title"].get()); dialog.destroy(); self.show_view("messages")

        ttk.Button(dialog, text=t("save", self.app.lang), command=save).grid(row=len(fields), column=0, columnspan=2, pady=12)


class StudentCabinetApp(Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title(APP_TITLE)
        self.geometry("1180x760")
        self.minsize(900, 600)
        self.lang = "ru"
        self.db = DatabaseManager()
        self.conn = self.db.conn
        self.current_user: dict[str, Any] | None = None
        self.login_window: LoginWindow | None = None
        self.dashboard: BaseDashboard | None = None
        self.show_login()

    def set_language(self, lang: str) -> None:
        self.lang = lang if lang in {"ru", "tg"} else "ru"
        if self.login_window is not None:
            self.login_window.status.configure(text="")
        if self.dashboard is not None:
            self.dashboard.show_view("dashboard")

    def show_login(self) -> None:
        if self.dashboard is not None:
            self.dashboard.destroy(); self.dashboard = None
        if self.login_window is not None:
            self.login_window.destroy()
        self.login_window = LoginWindow(self, self)

    def login(self, username: str, password: str) -> None:
        row = self.conn.execute("SELECT * FROM users WHERE username = ? AND is_active = 1", (username,)).fetchone()
        if row is None or not AuthenticationService.verify_password(row["password_hash"], password):
            raise ValueError(t("invalid_login", self.lang))
        self.current_user = {"id": row["id"], "username": row["username"], "role": row["role"]}
        if self.login_window is not None:
            self.login_window.destroy(); self.login_window = None
        dashboard_cls = AdminDashboard if row["role"] == "admin" else StudentDashboard
        self.dashboard = dashboard_cls(self, self, self.current_user)

    def logout(self) -> None:
        self.current_user = None
        if self.dashboard is not None:
            self.dashboard.destroy(); self.dashboard = None
        self.show_login()


def run_client() -> None:
    app = StudentCabinetApp()
    app.mainloop()


def main() -> None:
    parser = argparse.ArgumentParser(description="Student Cabinet desktop application")
    parser.add_argument("mode", nargs="?", choices=["client", "server"], default="client")
    args = parser.parse_args()
    if args.mode == "server":
        print("This build is desktop-only. Run: python main.py client")
        return
    run_client()


if __name__ == "__main__":
    main()
'''

(root / 'main.py').write_text(main_code, encoding='utf-8')
(root / 'requirements.txt').write_text('argon2-cffi==25.1.0\nargon2-cffi-bindings==26.1.0\net_xmlfile==2.0.0\nopenpyxl==3.1.5\n', encoding='utf-8')
(root / 'README.md').write_text('''# Student Cabinet\n\nStudent Cabinet is a local desktop application for a single Windows workstation. It uses Python tkinter, SQLite, Argon2 password hashing, and openpyxl for Excel export. No PostgreSQL or FastAPI server is required.\n\n## Install\n\n```powershell\npy -3.11 -m venv .venv\n.\\.venv\\Scripts\\python.exe -m pip install -r requirements.txt\n.\\.venv\\Scripts\\python.exe main.py\n```\n\n## Default accounts\n\n- Admin: admin / admin123456\n- Student: student1 / student123\n\n## Features\n\n- Admin and student roles\n- Student profile, grades, schedule, attendance, payments and messages\n- Local SQLite database\n- Excel export\n- Russian and Tajik language switch\n- Password hashing via Argon2\n\n## Backup\n\nThe database is kept in `student_cabinet.db` next to the project files. Copy it to make a backup.\n''', encoding='utf-8')
(root / '.env').write_text('APP_NAME=Student Cabinet\nAPP_LANGUAGE=ru\nDB_PATH=student_cabinet.db\nDEFAULT_ADMIN_USERNAME=admin\nDEFAULT_ADMIN_PASSWORD=admin123456\nDEFAULT_STUDENT_USERNAME=student1\nDEFAULT_STUDENT_PASSWORD=student123\n', encoding='utf-8')
(root / '.env.example').write_text('APP_NAME=Student Cabinet\nAPP_LANGUAGE=ru\nDB_PATH=student_cabinet.db\nDEFAULT_ADMIN_USERNAME=admin\nDEFAULT_ADMIN_PASSWORD=admin123456\nDEFAULT_STUDENT_USERNAME=student1\nDEFAULT_STUDENT_PASSWORD=student123\n', encoding='utf-8')
print('updated main.py and local config files')
