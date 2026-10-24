from __future__ import annotations

import argparse
import json
import logging
import os
import sqlite3
import sys
import tkinter as tk
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

import customtkinter as ctk
from argon2 import PasswordHasher
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from tkinter import filedialog, messagebox, ttk


DB_PATH = Path(__file__).with_name("student_cabinet.db")
APP_TITLE = "Кабинет студента"
CURRENT_LANG = "ru"
APP_IMAGE_PATH = Path(__file__).with_name("assets") / "campus_illustration.ppm"

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger("student_cabinet")

TRANSLATIONS = {
    "ru": {
        "app_name": "Кабинет студента",
        "login": "Вход в систему",
        "username": "Логин",
        "password": "Пароль",
        "signin": "Войти",
        "logout": "Выйти",
        "dashboard": "Главная",
        "profile": "Профиль",
        "admin_profile_title": "Профиль администратора",
        "account_id": "Код учётной записи",
        "account_status": "Статус учётной записи",
        "created_at": "Дата создания",
        "grades": "Оценки",
        "schedule": "Расписание",
        "attendance": "Посещаемость",
        "payments": "Оплата",
        "messages": "Сообщения",
        "students": "Студенты",
        "groups": "Группы",
        "admin": "Администратор",
        "student": "Студент",
        "search": "Поиск",
        "export": "Экспорт в Excel",
        "no_data": "Нет данных",
        "save": "Сохранить",
        "cancel": "Отмена",
        "error": "Ошибка",
        "success": "Успешно",
        "apply": "Применить",
        "password_changed": "Пароль изменён",
        "welcome": "Добро пожаловать",
        "add": "Добавить",
        "edit": "Изменить",
        "delete": "Удалить",
        "mark_read": "Отметить прочитанным",
        "current_password": "Текущий пароль",
        "new_password": "Новый пароль",
        "app_subtitle": "Единое пространство для учёбы и управления",
        "login_hint": "Введите данные своей учётной записи",
        "role_admin": "Администратор",
        "role_student": "Студент",
        "admin_greeting": "Здравствуйте",
        "student_greeting": "Здравствуйте",
        "today": "Сегодня",
        "overview": "Краткий обзор",
        "total_students": "Студенты",
        "total_groups": "Группы",
        "total_subjects": "Предметы",
        "avg_grade": "Средний балл",
        "attendance_rate": "Посещаемость",
        "schedule_count": "Занятия в расписании",
        "balance_due": "Остаток к оплате",
        "recent_messages": "Последние объявления",
        "quick_actions": "Быстрый переход",
        "open_section": "Открыть раздел",
        "all_students": "Список студентов",
        "catalogs": "Справочники",
        "add_student": "Добавить студента",
        "add_grade": "Внести оценку",
        "add_schedule": "Добавить занятие",
        "add_attendance": "Отметить посещаемость",
        "add_payment": "Записать оплату",
        "add_message": "Создать объявление",
        "export": "Экспорт в Excel",
        "no_data": "Пока нет данных",
        "no_student_profile": "Профиль студента не найден. Обратитесь к администратору.",
        "show_password": "Показать",
        "hide_password": "Скрыть",
        "login_error": "Неверный логин или пароль",
        "required_fields": "Заполните все обязательные поля",
        "invalid_value": "Проверьте введённые данные",
        "current_password_error": "Неверно указан текущий пароль",
        "password_min_length": "Новый пароль должен содержать не менее 8 символов",
        "form_title_student": "Карточка студента",
        "form_title_grade": "Новая оценка",
        "form_title_schedule": "Новое занятие",
        "form_title_attendance": "Посещаемость",
        "form_title_payment": "Новая оплата",
        "form_title_message": "Новое объявление",
        "student_number": "Номер студента",
        "full_name": "ФИО",
        "group_id": "Код группы",
        "specialty_id": "Код специальности",
        "course_id": "Код курса",
        "email": "Электронная почта",
        "phone": "Телефон",
        "birth_date": "Дата рождения (ГГГГ-ММ-ДД)",
        "student_id": "Код студента",
        "subject_id": "Код предмета",
        "semester_id": "Код семестра",
        "current_score": "Текущий балл",
        "examination_score": "Экзаменационный балл",
        "final_score": "Итоговый балл",
        "note": "Примечание",
        "weekday": "День недели (0–6)",
        "starts_at": "Начало (ЧЧ:ММ)",
        "ends_at": "Окончание (ЧЧ:ММ)",
        "teacher": "Преподаватель",
        "classroom": "Аудитория",
        "lesson_date": "Дата занятия (ГГГГ-ММ-ДД)",
        "status": "Статус посещаемости",
        "notes": "Примечания",
        "starts_on": "Дата начала",
        "ends_on": "Дата окончания",
        "year_number": "Номер курса",
        "term": "Семестр (1 или 2)",
        "contract_id": "Код договора",
        "amount": "Сумма",
        "payment_date": "Дата оплаты (ГГГГ-ММ-ДД)",
        "receipt_number": "Номер квитанции",
        "target_group_id": "Код группы (пусто — всем)",
        "title": "Заголовок",
        "body": "Текст объявления",
        "group_name": "Группа",
        "course_name": "Курс",
        "specialty_name": "Специальность",
        "subject_name": "Предмет",
        "semester_name": "Семестр",
        "academic_year": "Учебный год",
        "contract_number": "Номер договора",
        "contract_id": "Код договора",
        "outstanding_balance": "Остаток",
        "currency": "Валюта",
        "published_at": "Дата публикации",
        "is_published": "Опубликовано",
        "read": "Прочитано",
        "payments": "Оплаты",
        "view_profile": "Профиль",
        "admin_welcome": "Рабочее пространство администратора",
        "student_welcome": "Ваш учебный прогресс в одном месте",
        "loading_error": "Не удалось загрузить данные",
        "attendance_present": "Присутствовал",
        "attendance_absent": "Отсутствовал",
        "attendance_excused": "Уважительная причина",
        "language": "Язык",
        "logout_confirm": "Выйти из учётной записи?",
        "active": "Активен",
        "inactive": "Неактивен",
        "yes": "Да",
        "no": "Нет",
        "catalog_groups": "Учебные группы",
        "catalog_subjects": "Предметы",
        "catalog_specialties": "Специальности",
        "catalog_courses": "Курсы",
        "catalog_semesters": "Семестры",
        "catalog_contracts": "Договоры",
        "excel_files": "Файлы Excel",
        "contract_created": "Договор добавлен",
        "select_none": "Не выбрано",
        "report": "Отчёт",
        "id": "Код",
        "name": "Название",
        "code": "Код предмета",
        "is_active": "Статус учётной записи",
        "toggle_student": "Изменить статус",
        "edit_student": "Изменить карточку",
        "student_selected": "Сначала выберите студента в таблице",
        "student_deactivated": "Учётная запись студента отключена",
        "student_activated": "Учётная запись студента активирована",
        "payment_exceeds_balance": "Сумма превышает остаток по договору",
        "grade_weights_missing": "Укажите итоговый балл или настройте веса оценивания",
        "catalogs_hint": "Сначала заполните справочники, затем добавляйте студентов и учебные записи.",
    },
    "tg": {
        "app_name": "Кабинети донишҷӯ",
        "login": "Воридшавӣ ба система",
        "username": "Номи корбар",
        "password": "Рамз",
        "signin": "Даромадан",
        "logout": "Баромадан",
        "dashboard": "Асосӣ",
        "profile": "Профил",
        "admin_profile_title": "Профили маъмур",
        "account_id": "Рамзи ҳисоб",
        "account_status": "Вазъияти ҳисоб",
        "created_at": "Санаи таъсис",
        "grades": "Баҳоҳо",
        "schedule": "Ҷадвал",
        "attendance": "Иштирок",
        "payments": "Пардохт",
        "messages": "Паёмҳо",
        "students": "Донишҷӯён",
        "groups": "Гурӯҳҳо",
        "admin": "Маъмур",
        "student": "Донишҷӯ",
        "search": "Ҷустуҷӯ",
        "export": "Содирот ба Excel",
        "no_data": "Маълумоти нест",
        "save": "Захира",
        "cancel": "Бекор",
        "error": "Хато",
        "success": "Муваффақият",
        "apply": "Иҷро",
        "password_changed": "Рамз иваз шуд",
        "welcome": "Хуш омадед",
        "add": "Илова",
        "edit": "Тағир",
        "delete": "Нест",
        "mark_read": "Хондаро қайд кардан",
        "current_password": "Рамзи ҷорӣ",
        "new_password": "Рамзи нав",
        "app_subtitle": "Муҳити ягона барои таҳсил ва идоракунӣ",
        "login_hint": "Маълумоти ҳисоби худро ворид кунед",
        "role_admin": "Маъмур",
        "role_student": "Донишҷӯ",
        "admin_greeting": "Салом",
        "student_greeting": "Салом",
        "today": "Имрӯз",
        "overview": "Шарҳи кӯтоҳ",
        "total_students": "Донишҷӯён",
        "total_groups": "Гурӯҳҳо",
        "total_subjects": "Фанҳо",
        "avg_grade": "Баҳои миёна",
        "attendance_rate": "Иштирок",
        "schedule_count": "Дарсҳо дар ҷадвал",
        "balance_due": "Бақияи пардохт",
        "recent_messages": "Эълонҳои охирин",
        "quick_actions": "Гузариши зуд",
        "open_section": "Кушодани бахш",
        "all_students": "Рӯйхати донишҷӯён",
        "catalogs": "Маълумотномаҳо",
        "add_student": "Иловаи донишҷӯ",
        "add_grade": "Гузоштани баҳо",
        "add_schedule": "Иловаи дарс",
        "add_attendance": "Қайд кардани иштирок",
        "add_payment": "Сабти пардохт",
        "add_message": "Эҷоди эълон",
        "export": "Содирот ба Excel",
        "no_data": "Ҳоло маълумот нест",
        "no_student_profile": "Профили донишҷӯ ёфт нашуд. Ба маъмур муроҷиат кунед.",
        "show_password": "Намоиш",
        "hide_password": "Пинҳон",
        "login_error": "Номи корбар ё рамз нодуруст аст",
        "required_fields": "Ҳамаи майдонҳои ҳатмиро пур кунед",
        "invalid_value": "Маълумоти воридшударо санҷед",
        "current_password_error": "Рамзи ҷорӣ нодуруст аст",
        "password_min_length": "Рамзи нав бояд камаш 8 аломат дошта бошад",
        "form_title_student": "Маълумоти донишҷӯ",
        "form_title_grade": "Баҳои нав",
        "form_title_schedule": "Дарси нав",
        "form_title_attendance": "Иштирок",
        "form_title_payment": "Пардохти нав",
        "form_title_message": "Эълони нав",
        "student_number": "Рақами донишҷӯ",
        "full_name": "Ному насаб",
        "group_id": "Рамзи гурӯҳ",
        "specialty_id": "Рамзи ихтисос",
        "course_id": "Рамзи курс",
        "email": "Почтаи электронӣ",
        "phone": "Телефон",
        "birth_date": "Санаи таваллуд (СОЛ-МОҲ-РӮЗ)",
        "student_id": "Рамзи донишҷӯ",
        "subject_id": "Рамзи фан",
        "semester_id": "Рамзи нимсола",
        "current_score": "Баҳои ҷорӣ",
        "examination_score": "Баҳои имтиҳон",
        "final_score": "Баҳои ниҳоӣ",
        "note": "Шарҳ",
        "weekday": "Рӯзи ҳафта (0–6)",
        "starts_at": "Оғоз (СС:ДД)",
        "ends_at": "Анҷом (СС:ДД)",
        "teacher": "Омӯзгор",
        "classroom": "Синфхона",
        "lesson_date": "Санаи дарс (СОЛ-МОҲ-РӮЗ)",
        "status": "Вазъияти иштирок",
        "notes": "Шарҳҳо",
        "starts_on": "Санаи оғоз",
        "ends_on": "Санаи анҷом",
        "year_number": "Рақами курс",
        "term": "Нимсола (1 ё 2)",
        "contract_id": "Рамзи шартнома",
        "amount": "Маблағ",
        "payment_date": "Санаи пардохт (СОЛ-МОҲ-РӮЗ)",
        "receipt_number": "Рақами расид",
        "target_group_id": "Рамзи гурӯҳ (холӣ — ҳама)",
        "title": "Сарлавҳа",
        "body": "Матни эълон",
        "group_name": "Гурӯҳ",
        "course_name": "Курс",
        "specialty_name": "Ихтисос",
        "subject_name": "Фан",
        "semester_name": "Нимсола",
        "academic_year": "Соли таҳсил",
        "contract_number": "Рақами шартнома",
        "contract_id": "Рамзи шартнома",
        "outstanding_balance": "Бақия",
        "currency": "Асъор",
        "published_at": "Санаи нашр",
        "is_published": "Нашр шуд",
        "read": "Хонда шуд",
        "payments": "Пардохтҳо",
        "view_profile": "Профил",
        "admin_welcome": "Муҳити кории маъмур",
        "student_welcome": "Пешрафти таҳсили шумо дар як ҷо",
        "loading_error": "Маълумотро бор кардан муяссар нашуд",
        "attendance_present": "Ҳозир",
        "attendance_absent": "Ғоиб",
        "attendance_excused": "Сабаби узрнок",
        "language": "Забон",
        "logout_confirm": "Аз ҳисоб баромадан?",
        "active": "Фаъол",
        "inactive": "Ғайрифаъол",
        "yes": "Ҳа",
        "no": "Не",
        "catalog_groups": "Гурӯҳҳои таълимӣ",
        "catalog_subjects": "Фанҳо",
        "catalog_specialties": "Ихтисосҳо",
        "catalog_courses": "Курсҳо",
        "catalog_semesters": "Нимсолаҳо",
        "catalog_contracts": "Шартномаҳо",
        "excel_files": "Файлҳои Excel",
        "contract_created": "Шартнома илова шуд",
        "select_none": "Интихоб нашудааст",
        "report": "Ҳисобот",
        "id": "Рамз",
        "name": "Ном",
        "code": "Рамзи фан",
        "is_active": "Вазъияти ҳисоб",
        "toggle_student": "Тағйир додани вазъият",
        "edit_student": "Тағйири маълумоти донишҷӯ",
        "student_selected": "Аввал донишҷӯро аз ҷадвал интихоб кунед",
        "student_deactivated": "Ҳисоби донишҷӯ ғайрифаъол карда шуд",
        "student_activated": "Ҳисоби донишҷӯ фаъол карда шуд",
        "payment_exceeds_balance": "Маблағ аз бақияи шартнома зиёд аст",
        "grade_weights_missing": "Баҳои ниҳоиро ворид кунед ё вазни баҳогузориро танзим намоед",
        "catalogs_hint": "Аввал маълумотномаҳоро пур кунед, сипас донишҷӯён ва сабтҳои таълимиро илова намоед.",
    },
}


class DatabaseManager:
    """Local SQLite database manager for a single-computer Student Cabinet."""

    def __init__(self, db_path: str | Path | None = None) -> None:
        self.db_path = Path(db_path) if db_path else DB_PATH
        self.conn: sqlite3.Connection | None = None
        self.connect()

    def connect(self) -> None:
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self.conn = sqlite3.connect(str(self.db_path))
        self.conn.row_factory = sqlite3.Row
        self.conn.execute("PRAGMA foreign_keys = ON")
        self.init_schema()
        self.ensure_bootstrap_admin()

    def init_schema(self) -> None:
        if self.conn is None:
            raise RuntimeError("Database connection not initialized.")
        self.conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT NOT NULL UNIQUE,
                password_hash TEXT NOT NULL,
                role TEXT NOT NULL CHECK(role IN ('student','admin')),
                is_active INTEGER NOT NULL DEFAULT 1,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );

            CREATE TABLE IF NOT EXISTS groups (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE
            );

            CREATE TABLE IF NOT EXISTS specialties (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE
            );

            CREATE TABLE IF NOT EXISTS courses (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE,
                year_number INTEGER NOT NULL UNIQUE
            );

            CREATE TABLE IF NOT EXISTS semesters (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                academic_year TEXT NOT NULL,
                term INTEGER NOT NULL,
                starts_on TEXT NOT NULL,
                ends_on TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS students (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL UNIQUE,
                student_number TEXT NOT NULL UNIQUE,
                full_name TEXT NOT NULL,
                group_id INTEGER NOT NULL,
                specialty_id INTEGER,
                course_id INTEGER,
                email TEXT,
                phone TEXT,
                birth_date TEXT,
                FOREIGN KEY(user_id) REFERENCES users(id),
                FOREIGN KEY(group_id) REFERENCES groups(id),
                FOREIGN KEY(specialty_id) REFERENCES specialties(id),
                FOREIGN KEY(course_id) REFERENCES courses(id)
            );

            CREATE TABLE IF NOT EXISTS subjects (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE,
                code TEXT UNIQUE
            );

            CREATE TABLE IF NOT EXISTS enrollments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL,
                subject_id INTEGER NOT NULL,
                semester_id INTEGER NOT NULL,
                UNIQUE(student_id, subject_id, semester_id),
                FOREIGN KEY(student_id) REFERENCES students(id),
                FOREIGN KEY(subject_id) REFERENCES subjects(id),
                FOREIGN KEY(semester_id) REFERENCES semesters(id)
            );

            CREATE TABLE IF NOT EXISTS grades (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL,
                subject_id INTEGER NOT NULL,
                semester_id INTEGER NOT NULL,
                current_score REAL,
                examination_score REAL,
                final_score REAL,
                note TEXT,
                updated_at TEXT DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(student_id, subject_id, semester_id),
                FOREIGN KEY(student_id) REFERENCES students(id),
                FOREIGN KEY(subject_id) REFERENCES subjects(id),
                FOREIGN KEY(semester_id) REFERENCES semesters(id)
            );

            CREATE TABLE IF NOT EXISTS schedule (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                group_id INTEGER NOT NULL,
                subject_id INTEGER NOT NULL,
                weekday INTEGER NOT NULL CHECK(weekday BETWEEN 0 AND 6),
                starts_at TEXT NOT NULL,
                ends_at TEXT NOT NULL,
                teacher TEXT NOT NULL,
                classroom TEXT NOT NULL,
                semester_id INTEGER,
                FOREIGN KEY(group_id) REFERENCES groups(id),
                FOREIGN KEY(subject_id) REFERENCES subjects(id),
                FOREIGN KEY(semester_id) REFERENCES semesters(id)
            );

            CREATE TABLE IF NOT EXISTS attendance (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL,
                subject_id INTEGER NOT NULL,
                lesson_date TEXT NOT NULL,
                status TEXT NOT NULL CHECK(status IN ('present','absent','excused')),
                notes TEXT,
                UNIQUE(student_id, subject_id, lesson_date),
                FOREIGN KEY(student_id) REFERENCES students(id),
                FOREIGN KEY(subject_id) REFERENCES subjects(id)
            );

            CREATE TABLE IF NOT EXISTS contracts (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL,
                amount REAL NOT NULL CHECK(amount >= 0),
                currency TEXT NOT NULL DEFAULT 'TJS',
                contract_number TEXT NOT NULL UNIQUE,
                academic_year TEXT NOT NULL,
                FOREIGN KEY(student_id) REFERENCES students(id)
            );

            CREATE TABLE IF NOT EXISTS payments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                contract_id INTEGER NOT NULL,
                amount REAL NOT NULL CHECK(amount > 0),
                payment_date TEXT NOT NULL,
                receipt_number TEXT NOT NULL UNIQUE,
                note TEXT,
                FOREIGN KEY(contract_id) REFERENCES contracts(id)
            );

            CREATE TABLE IF NOT EXISTS messages (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                author_id INTEGER NOT NULL,
                target_group_id INTEGER,
                title TEXT NOT NULL,
                body TEXT NOT NULL,
                published_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                is_published INTEGER NOT NULL DEFAULT 1,
                FOREIGN KEY(author_id) REFERENCES users(id),
                FOREIGN KEY(target_group_id) REFERENCES groups(id)
            );

            CREATE TABLE IF NOT EXISTS message_reads (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                message_id INTEGER NOT NULL,
                user_id INTEGER NOT NULL,
                read_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(message_id, user_id),
                FOREIGN KEY(message_id) REFERENCES messages(id),
                FOREIGN KEY(user_id) REFERENCES users(id)
            );

            CREATE TABLE IF NOT EXISTS audit_logs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                actor_user_id INTEGER,
                action TEXT NOT NULL,
                entity_type TEXT NOT NULL,
                entity_id TEXT,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY(actor_user_id) REFERENCES users(id)
            );
            """
        )
        self.conn.commit()

    def ensure_bootstrap_admin(self) -> None:
        if self.conn is None:
            raise RuntimeError("Database connection not initialized.")
        admin_row = self.conn.execute("SELECT COUNT(*) AS count FROM users WHERE role = 'admin'").fetchone()
        if not admin_row or admin_row["count"] == 0:
            admin_username = os.getenv("INITIAL_ADMIN_USERNAME", "admin").strip() or "admin"
            admin_password = os.getenv("INITIAL_ADMIN_PASSWORD", "admin123456").strip() or "admin123456"
            self.conn.execute(
                "INSERT INTO users (username, password_hash, role, is_active) VALUES (?, ?, 'admin', 1)",
                (admin_username, AuthenticationService.hash_password(admin_password)),
            )
            self.conn.commit()
            logger.warning("Initial admin created using environment or fallback credentials.")

    def fetch_all(self, query: str, params: Iterable[Any] = ()) -> list[sqlite3.Row]:
        if self.conn is None:
            raise RuntimeError("Database connection not initialized.")
        return self.conn.execute(query, tuple(params)).fetchall()

    def execute(self, query: str, params: Iterable[Any] = ()) -> sqlite3.Cursor:
        if self.conn is None:
            raise RuntimeError("Database connection not initialized.")
        cursor = self.conn.execute(query, tuple(params))
        self.conn.commit()
        return cursor

    def close(self) -> None:
        if self.conn is not None:
            self.conn.close()
            self.conn = None


class AuthenticationService:
    """Password hashing and local plaintext-free auth."""

    _hasher = PasswordHasher()

    @staticmethod
    def hash_password(password: str) -> str:
        if len(password) < 8:
            raise ValueError("Password must be at least 8 characters long.")
        return AuthenticationService._hasher.hash(password)

    @staticmethod
    def verify_password(password_hash: str, password: str) -> bool:
        try:
            return AuthenticationService._hasher.verify(password_hash, password)
        except Exception:
            return False

    @staticmethod
    def create_session(user: dict[str, Any]) -> dict[str, Any]:
        return {"id": user["id"], "username": user["username"], "role": user["role"]}


class StudentService:
    """Student profile and account methods."""

    @staticmethod
    def get_by_user_id(conn: sqlite3.Connection, user_id: int) -> dict[str, Any] | None:
        row = conn.execute(
            "SELECT s.*, u.username, u.role, g.name AS group_name, "
            "sp.name AS specialty_name, c.name AS course_name "
            "FROM students s JOIN users u ON u.id = s.user_id "
            "LEFT JOIN groups g ON g.id = s.group_id "
            "LEFT JOIN specialties sp ON sp.id = s.specialty_id "
            "LEFT JOIN courses c ON c.id = s.course_id WHERE u.id = ?",
            (user_id,),
        ).fetchone()
        return dict(row) if row else None

    @staticmethod
    def list(conn: sqlite3.Connection, query_text: str | None = None) -> list[dict[str, Any]]:
        sql = (
            "SELECT s.*, u.username, u.is_active, g.name AS group_name, "
            "sp.name AS specialty_name, c.name AS course_name "
            "FROM students s JOIN users u ON u.id = s.user_id LEFT JOIN groups g ON g.id = s.group_id "
            "LEFT JOIN specialties sp ON sp.id = s.specialty_id LEFT JOIN courses c ON c.id = s.course_id"
        )
        params: list[Any] = []
        if query_text:
            like = f"%{query_text.strip()}%"
            sql += " WHERE s.full_name LIKE ? OR s.student_number LIKE ? OR u.username LIKE ?"
            params.extend([like, like, like])
        sql += " ORDER BY s.full_name"
        rows = conn.execute(sql, params).fetchall()
        return [dict(row) for row in rows]

    @staticmethod
    def create_student(conn: sqlite3.Connection, payload: dict[str, Any]) -> dict[str, Any]:
        user = conn.execute(
            "INSERT INTO users(username, password_hash, role, is_active) VALUES (?, ?, 'student', 1)",
            (payload["username"], AuthenticationService.hash_password(payload["password"])),
        )
        user_id = user.lastrowid
        student_id = conn.execute(
            "INSERT INTO students(user_id, student_number, full_name, group_id, specialty_id, course_id, email, phone, birth_date) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (
                user_id,
                payload.get("student_number"),
                payload.get("full_name"),
                payload.get("group_id"),
                payload.get("specialty_id"),
                payload.get("course_id"),
                payload.get("email"),
                payload.get("phone"),
                payload.get("birth_date"),
            ),
        ).lastrowid
        conn.commit()
        return {"id": student_id, "user_id": user_id, "username": payload["username"]}

    @staticmethod
    def update_student(conn: sqlite3.Connection, student_id: int, payload: dict[str, Any]) -> None:
        assignments = []
        params: list[Any] = []
        for key in [
            "student_number",
            "full_name",
            "group_id",
            "specialty_id",
            "course_id",
            "email",
            "phone",
            "birth_date",
        ]:
            if key in payload and payload[key] is not None:
                assignments.append(f"{key} = ?")
                params.append(payload[key])
        if assignments:
            conn.execute(f"UPDATE students SET {', '.join(assignments)} WHERE id = ?", params + [student_id])
            conn.commit()

    @staticmethod
    def deactivate_student(conn: sqlite3.Connection, student_id: int) -> None:
        user_id = conn.execute("SELECT user_id FROM students WHERE id = ?", (student_id,)).fetchone()
        if user_id is None:
            raise ValueError("Student not found")
        conn.execute("UPDATE users SET is_active = 0 WHERE id = ?", (user_id[0],))
        conn.commit()

    @staticmethod
    def activate_student(conn: sqlite3.Connection, student_id: int) -> None:
        user_id = conn.execute("SELECT user_id FROM students WHERE id = ?", (student_id,)).fetchone()
        if user_id is None:
            raise ValueError("Student not found")
        conn.execute("UPDATE users SET is_active = 1 WHERE id = ?", (user_id[0],))
        conn.commit()


class GradeService:
    """Grade management with configurable formula support."""

    @staticmethod
    def final_score(current_score: float, exam_score: float, configured_final: float | None = None) -> float:
        if configured_final is not None:
            return float(configured_final)
        current_weight = os.getenv("GRADE_CURRENT_WEIGHT")
        exam_weight = os.getenv("GRADE_EXAM_WEIGHT")
        if current_weight and exam_weight:
            return round(float(current_weight) * float(current_score) + float(exam_weight) * float(exam_score), 2)
        raise ValueError("Please enter final score or configure official grade weights.")

    @staticmethod
    def upsert(conn: sqlite3.Connection, payload: dict[str, Any]) -> dict[str, Any]:
        final_value = GradeService.final_score(
            float(payload["current_score"]),
            float(payload["examination_score"]),
            payload.get("final_score"),
        )
        row = conn.execute(
            "SELECT id FROM grades WHERE student_id = ? AND subject_id = ? AND semester_id = ?",
            (payload["student_id"], payload["subject_id"], payload["semester_id"]),
        ).fetchone()
        if row is None:
            cursor = conn.execute(
                "INSERT INTO grades(student_id, subject_id, semester_id, current_score, examination_score, final_score, note, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)",
                (
                    payload["student_id"],
                    payload["subject_id"],
                    payload["semester_id"],
                    payload["current_score"],
                    payload["examination_score"],
                    final_value,
                    payload.get("note"),
                ),
            )
            gid = cursor.lastrowid
        else:
            conn.execute(
                "UPDATE grades SET current_score = ?, examination_score = ?, final_score = ?, note = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                (
                    payload["current_score"],
                    payload["examination_score"],
                    final_value,
                    payload.get("note"),
                    row["id"],
                ),
            )
            gid = row["id"]
        conn.commit()
        return {"id": gid, "final_score": final_value}

    @staticmethod
    def list_for_student(conn: sqlite3.Connection, student_id: int) -> list[dict[str, Any]]:
        rows = conn.execute(
            "SELECT g.*, s.name AS subject_name, sem.name AS semester_name FROM grades g JOIN subjects s ON s.id = g.subject_id JOIN semesters sem ON sem.id = g.semester_id WHERE g.student_id = ? ORDER BY sem.academic_year DESC, sem.term, s.name",
            (student_id,),
        ).fetchall()
        return [dict(row) for row in rows]

    @staticmethod
    def list_all(conn: sqlite3.Connection) -> list[dict[str, Any]]:
        rows = conn.execute(
            "SELECT g.*, s.name AS subject_name, sem.name AS semester_name, st.full_name, st.student_number FROM grades g JOIN subjects s ON s.id = g.subject_id JOIN semesters sem ON sem.id = g.semester_id JOIN students st ON st.id = g.student_id ORDER BY st.full_name, s.name"
        ).fetchall()
        return [dict(row) for row in rows]


class ScheduleService:
    """Schedule management."""

    @staticmethod
    def list_for_group(conn: sqlite3.Connection, group_id: int) -> list[dict[str, Any]]:
        rows = conn.execute(
            "SELECT sc.*, subj.name AS subject_name FROM schedule sc JOIN subjects subj ON subj.id = sc.subject_id WHERE sc.group_id = ? ORDER BY sc.weekday, sc.starts_at",
            (group_id,),
        ).fetchall()
        return [dict(r) for r in rows]

    @staticmethod
    def list_all(conn: sqlite3.Connection) -> list[dict[str, Any]]:
        rows = conn.execute(
            "SELECT sc.*, subj.name AS subject_name, g.name AS group_name FROM schedule sc JOIN subjects subj ON subj.id = sc.subject_id JOIN groups g ON g.id = sc.group_id ORDER BY g.name, sc.weekday, sc.starts_at"
        ).fetchall()
        return [dict(r) for r in rows]


class AttendanceService:
    """Attendance registration."""

    @staticmethod
    def list_for_student(conn: sqlite3.Connection, student_id: int) -> list[dict[str, Any]]:
        rows = conn.execute(
            "SELECT a.*, s.name AS subject_name FROM attendance a JOIN subjects s ON s.id = a.subject_id WHERE a.student_id = ? ORDER BY a.lesson_date DESC",
            (student_id,),
        ).fetchall()
        return [dict(r) for r in rows]

    @staticmethod
    def percentage(records: list[dict[str, Any]]) -> float | None:
        if not records:
            return None
        valid = [r for r in records if r.get("status") != "excused"]
        if not valid:
            return None
        present = sum(1 for r in valid if r.get("status") == "present")
        return round((present / len(valid)) * 100, 2)


class PaymentService:
    """Local contract and payment tracking."""

    @staticmethod
    def outstanding_balance(conn: sqlite3.Connection, contract_id: int) -> float:
        contract = conn.execute("SELECT amount FROM contracts WHERE id = ?", (contract_id,)).fetchone()
        if contract is None:
            return 0.0
        paid = conn.execute("SELECT COALESCE(SUM(amount),0) FROM payments WHERE contract_id = ?", (contract_id,)).fetchone()[0]
        return float(contract["amount"]) - float(paid)

    @staticmethod
    def list_for_student(conn: sqlite3.Connection, student_id: int) -> list[dict[str, Any]]:
        contracts = conn.execute(
            "SELECT * FROM contracts WHERE student_id = ? ORDER BY academic_year DESC",
            (student_id,),
        ).fetchall()
        result = []
        for contract in contracts:
            rows = conn.execute(
                "SELECT * FROM payments WHERE contract_id = ? ORDER BY payment_date DESC",
                (contract["id"],),
            ).fetchall()
            result.append({
                "contract_id": contract["id"],
                "contract_number": contract["contract_number"],
                "academic_year": contract["academic_year"],
                "amount": contract["amount"],
                "currency": contract["currency"],
                "outstanding_balance": PaymentService.outstanding_balance(conn, contract["id"]),
                "payments": [dict(r) for r in rows],
            })
        return result

    @staticmethod
    def create_payment(conn: sqlite3.Connection, payload: dict[str, Any]) -> dict[str, Any]:
        balance = PaymentService.outstanding_balance(conn, payload["contract_id"])
        if float(payload["amount"]) > balance:
            raise ValueError("Payment exceeds remaining balance.")
        cursor = conn.execute(
            "INSERT INTO payments(contract_id, amount, payment_date, receipt_number, note) VALUES (?, ?, ?, ?, ?)",
            (
                payload["contract_id"],
                float(payload["amount"]),
                payload["payment_date"],
                payload["receipt_number"],
                payload.get("note"),
            ),
        )
        conn.commit()
        return {"id": cursor.lastrowid, "amount": payload["amount"]}


class MessageService:
    """Announcements and group messages."""

    @staticmethod
    def list_for_student(conn: sqlite3.Connection, user_id: int, student: dict[str, Any]) -> list[dict[str, Any]]:
        rows = conn.execute(
            "SELECT m.*, mr.id AS read_id FROM messages m LEFT JOIN message_reads mr ON mr.message_id = m.id AND mr.user_id = ? WHERE m.is_published = 1 AND (m.target_group_id IS NULL OR m.target_group_id = ?) ORDER BY m.published_at DESC",
            (user_id, student["group_id"]),
        ).fetchall()
        return [{
            "id": row["id"],
            "title": row["title"],
            "body": row["body"],
            "published_at": row["published_at"],
            "read": row["read_id"] is not None,
        } for row in rows]

    @staticmethod
    def mark_read(conn: sqlite3.Connection, message_id: int, user_id: int) -> None:
        exists = conn.execute("SELECT 1 FROM message_reads WHERE message_id = ? AND user_id = ?", (message_id, user_id)).fetchone()
        if exists is None:
            conn.execute("INSERT INTO message_reads(message_id, user_id, read_at) VALUES (?, ?, CURRENT_TIMESTAMP)", (message_id, user_id))
            conn.commit()

    @staticmethod
    def list_all(conn: sqlite3.Connection) -> list[dict[str, Any]]:
        rows = conn.execute(
            "SELECT m.*, g.name AS group_name FROM messages m LEFT JOIN groups g ON g.id = m.target_group_id ORDER BY m.published_at DESC"
        ).fetchall()
        return [dict(r) for r in rows]


class ExcelExporter:
    """Export local data to Excel files."""

    @staticmethod
    def export(
        rows: list[dict[str, Any]],
        target: str | Path,
        sheet_title: str = "Report",
        empty_label: str = "No data",
    ) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_title
        if rows:
            headers = list(rows[0].keys())
            ws.append(headers)
            for row in rows:
                ws.append([row.get(h) for h in headers])
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="17365D")
        else:
            ws.append([empty_label])
        out = Path(target)
        out.parent.mkdir(parents=True, exist_ok=True)
        wb.save(out)
        return out


class APIClient:
    """Compatibility shim for the old API design. This app uses SQLite directly."""

    def __init__(self) -> None:
        self.db = DatabaseManager()

    def login(self, username: str, password: str) -> dict[str, Any]:
        user = self.db.conn.execute(
            "SELECT * FROM users WHERE username = ? AND is_active = 1",
            (username,),
        ).fetchone()
        if user is None or not AuthenticationService.verify_password(user["password_hash"], password):
            raise ValueError("Invalid username or password")
        return {"id": user["id"], "username": user["username"], "role": user["role"]}

    def get_student_profile(self, user_id: int) -> dict[str, Any]:
        student = StudentService.get_by_user_id(self.db.conn, user_id)
        if student is None:
            raise ValueError("Student not found")
        return student

    def get_student_messages(self, user_id: int) -> list[dict[str, Any]]:
        student = StudentService.get_by_user_id(self.db.conn, user_id)
        if student is None:
            return []
        return MessageService.list_for_student(self.db.conn, user_id, student)


class LoginWindow(ctk.CTkFrame):
    """Login view displayed inside the application root window."""

    def __init__(self, master: ctk.CTk, callback: Any, language: str = "ru") -> None:
        super().__init__(master)
        self.callback = callback
        self.language = language
        self.configure(fg_color="#F4F7FB")
        master.title(TRANSLATIONS[language]["app_name"])
        self.grid_columnconfigure(0, weight=6, uniform="login")
        self.grid_columnconfigure(1, weight=5, uniform="login")
        self.grid_rowconfigure(0, weight=1)

        t = TRANSLATIONS[language]
        hero = ctk.CTkFrame(self, corner_radius=0, fg_color="#102A4C")
        hero.grid(row=0, column=0, sticky="nsew")
        hero.grid_columnconfigure(0, weight=1)
        hero.grid_rowconfigure(3, weight=1)
        self.brand_label = ctk.CTkLabel(
            hero,
            text="✦  " + t["app_name"],
            font=ctk.CTkFont(size=24, weight="bold"),
            text_color="#FFFFFF",
        )
        self.brand_label.grid(row=0, column=0, padx=46, pady=(42, 22), sticky="w")
        self.subtitle_label = ctk.CTkLabel(
            hero,
            text=t["app_subtitle"],
            font=ctk.CTkFont(size=24, weight="bold"),
            text_color="#FFFFFF",
            wraplength=420,
            justify="left",
        )
        self.subtitle_label.grid(row=1, column=0, padx=46, sticky="w")
        self.welcome_label = ctk.CTkLabel(
            hero,
            text=t["student_welcome"],
            font=ctk.CTkFont(size=15),
            text_color="#C7D5E7",
            wraplength=420,
            justify="left",
        )
        self.welcome_label.grid(row=2, column=0, padx=46, pady=(12, 0), sticky="w")
        self.hero_image = tk.PhotoImage(file=str(APP_IMAGE_PATH))
        tk.Label(hero, image=self.hero_image, bg="#102A4C", bd=0).grid(
            row=3, column=0, padx=24, pady=18, sticky="se"
        )

        panel = ctk.CTkFrame(self, fg_color="#FFFFFF", corner_radius=20)
        panel.grid(row=0, column=1, padx=42, pady=46, sticky="nsew")
        panel.grid_columnconfigure(0, weight=1)
        panel.grid_rowconfigure(6, weight=1)
        self.login_title = ctk.CTkLabel(
            panel,
            text=t["login"],
            font=ctk.CTkFont(size=29, weight="bold"),
            text_color="#132E50",
        )
        self.login_title.grid(row=0, column=0, padx=38, pady=(54, 8), sticky="w")
        self.login_hint = ctk.CTkLabel(
            panel,
            text=t["login_hint"],
            font=ctk.CTkFont(size=14),
            text_color="#75849A",
        )
        self.login_hint.grid(row=1, column=0, padx=38, pady=(0, 26), sticky="w")
        self.username = ctk.CTkEntry(
            panel, height=48, placeholder_text=t["username"], corner_radius=10
        )
        self.username.grid(row=2, column=0, padx=38, pady=8, sticky="ew")
        password_row = ctk.CTkFrame(panel, fg_color="transparent")
        password_row.grid(row=3, column=0, padx=38, pady=8, sticky="ew")
        password_row.grid_columnconfigure(0, weight=1)
        self.password = ctk.CTkEntry(
            password_row, height=48, placeholder_text=t["password"], show="*",
            corner_radius=10,
        )
        self.password.grid(row=0, column=0, sticky="ew")
        self.password_visible = False
        self.password_toggle = ctk.CTkButton(
            password_row,
            width=76,
            height=42,
            text=t["show_password"],
            fg_color="transparent",
            text_color="#52647C",
            hover_color="#EEF2F7",
            command=self.toggle_password,
        )
        self.password_toggle.grid(row=0, column=1, padx=(6, 0))
        self.language_var = ctk.StringVar(
            value="Русский" if language == "ru" else "Тоҷикӣ"
        )
        self.lang_menu = ctk.CTkOptionMenu(
            panel,
            height=42,
            values=["Русский", "Тоҷикӣ"],
            variable=self.language_var,
            command=self.switch_language,
            fg_color="#EAF0F7",
            button_color="#D7E1ED",
            button_hover_color="#C6D5E5",
            text_color="#253A55",
        )
        self.lang_menu.grid(row=4, column=0, padx=38, pady=(14, 8), sticky="ew")
        self.signin = ctk.CTkButton(
            panel,
            height=48,
            corner_radius=10,
            text=t["signin"],
            command=self.try_login,
            fg_color="#176B87",
            hover_color="#12566D",
            font=ctk.CTkFont(size=15, weight="bold"),
        )
        self.signin.grid(row=5, column=0, padx=38, pady=(12, 8), sticky="ew")
        self.status = ctk.CTkLabel(
            panel, text="", text_color="#B42318", wraplength=340
        )
        self.status.grid(row=6, column=0, padx=38, pady=(0, 14), sticky="n")

        self.bind("<Return>", lambda _event: self.try_login())
        self.username.focus_set()

    def toggle_password(self) -> None:
        self.password_visible = not self.password_visible
        self.password.configure(show="" if self.password_visible else "*")
        self.password_toggle.configure(
            text=TRANSLATIONS[self.language][
                "hide_password" if self.password_visible else "show_password"
            ]
        )

    def switch_language(self, selection: str) -> None:
        self.language = "ru" if selection == "Русский" else "tg"
        translations = TRANSLATIONS[self.language]
        self.master.title(translations["app_name"])
        self.brand_label.configure(text="✦  " + translations["app_name"])
        self.subtitle_label.configure(text=translations["app_subtitle"])
        self.welcome_label.configure(text=translations["student_welcome"])
        self.login_title.configure(text=translations["login"])
        self.login_hint.configure(text=translations["login_hint"])
        self.username.configure(placeholder_text=translations["username"])
        self.password.configure(placeholder_text=translations["password"])
        self.signin.configure(text=translations["signin"])
        self.password_toggle.configure(
            text=translations[
                "hide_password" if self.password_visible else "show_password"
            ]
        )
        self.status.configure(text="")

    def try_login(self) -> None:
        username = self.username.get().strip()
        password = self.password.get()
        if not username or not password:
            self.status.configure(text=TRANSLATIONS[self.language]["required_fields"])
            return
        try:
            self.callback(username, password, self.language)
        except ValueError:
            self.status.configure(text=TRANSLATIONS[self.language]["login_error"])
            self.password.delete(0, "end")
        except Exception:
            logger.exception("Login failed unexpectedly")
            self.status.configure(
                text=TRANSLATIONS[self.language]["loading_error"]
            )
            self.password.delete(0, "end")


class StudentDashboard(ctk.CTkFrame):
    """Student dashboard with role-limited views."""

    def __init__(self, parent: Any, app: Any, user: dict[str, Any], language: str) -> None:
        super().__init__(parent, fg_color="#F3F5F8")
        self.app = app
        self.user = user
        self.language = language
        self.pack(fill="both", expand=True)
        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)
        self.build_sidebar()
        self.workspace = ctk.CTkFrame(self, fg_color="#F4F7FB", corner_radius=0)
        self.workspace.grid(row=0, column=1, sticky="nsew")
        self.workspace.grid_columnconfigure(0, weight=1)
        self.workspace.grid_rowconfigure(1, weight=1)
        self.header = ctk.CTkFrame(self.workspace, fg_color="#FFFFFF", corner_radius=0, height=76)
        self.header.grid(row=0, column=0, sticky="ew")
        self.header.grid_columnconfigure(0, weight=1)
        self.page_title = ctk.CTkLabel(
            self.header,
            text="",
            font=ctk.CTkFont(size=22, weight="bold"),
            text_color="#172B4D",
        )
        self.page_title.grid(row=0, column=0, padx=28, pady=20, sticky="w")
        self.header_user = ctk.CTkLabel(
            self.header,
            text=f"{self.user['username']}  ·  {self.t()['role_student']}",
            font=ctk.CTkFont(size=13, weight="bold"),
            text_color="#496078",
        )
        self.header_user.grid(row=0, column=1, padx=28, pady=20, sticky="e")
        self.content = ctk.CTkScrollableFrame(
            self.workspace, fg_color="#F4F7FB", corner_radius=0
        )
        self.content.grid(row=1, column=0, padx=22, pady=18, sticky="nsew")
        self.content.grid_columnconfigure(0, weight=1)
        self.show_view("dashboard")

    def t(self) -> dict[str, str]:
        return TRANSLATIONS[self.language]

    def build_sidebar(self) -> None:
        sidebar = ctk.CTkFrame(self, width=246, corner_radius=0, fg_color="#102A4C")
        sidebar.grid(row=0, column=0, sticky="nsew")
        sidebar.grid_propagate(False)
        sidebar.grid_rowconfigure(8, weight=1)
        ctk.CTkLabel(
            sidebar,
            text="✦  " + self.t()["app_name"],
            text_color="#FFFFFF",
            font=ctk.CTkFont(size=19, weight="bold"),
        ).grid(row=0, column=0, padx=20, pady=(26, 34), sticky="w")
        options = [
            ("dashboard", "⌂", self.t()["dashboard"]),
            ("profile", "◉", self.t()["profile"]),
            ("grades", "✧", self.t()["grades"]),
            ("schedule", "▦", self.t()["schedule"]),
            ("attendance", "✓", self.t()["attendance"]),
            ("payments", "¤", self.t()["payments"]),
            ("messages", "✉", self.t()["messages"]),
        ]
        self.nav_buttons: dict[str, ctk.CTkButton] = {}
        for row, (key, icon, label) in enumerate(options, start=1):
            button = ctk.CTkButton(
                sidebar,
                text=f"{icon}    {label}",
                command=lambda value=key: self.show_view(value),
                fg_color="#1D4268" if key == "dashboard" else "transparent",
                hover_color="#244C78",
                text_color="#E9F0F8",
                anchor="w",
                height=44,
                corner_radius=10,
                font=ctk.CTkFont(size=14),
            )
            button.grid(row=row, column=0, padx=12, pady=3, sticky="ew")
            self.nav_buttons[key] = button
        ctk.CTkLabel(
            sidebar,
            text=f"{self.user['username']}  ·  {self.t()['role_student']}",
            text_color="#B9C9DC",
            font=ctk.CTkFont(size=12),
        ).grid(row=9, column=0, padx=20, pady=(12, 10), sticky="w")
        ctk.CTkButton(
            sidebar,
            text=self.t()["password_changed"],
            command=self.change_password,
            fg_color="#1D4268",
            hover_color="#244C78",
            height=40,
            corner_radius=10,
        ).grid(row=10, column=0, padx=12, pady=4, sticky="ew")
        ctk.CTkButton(
            sidebar,
            text="↪   " + self.t()["logout"],
            command=self.app.logout,
            fg_color="#A64045",
            hover_color="#8F3339",
            height=42,
            corner_radius=10,
        ).grid(row=11, column=0, padx=12, pady=(4, 18), sticky="ew")

    def show_view(self, view: str) -> None:
        for child in self.content.winfo_children():
            child.destroy()
        self.data_table: ttk.Treeview | None = None
        for key, button in self.nav_buttons.items():
            button.configure(
                fg_color="#1D4268" if key == view else "transparent"
            )
        self.page_title.configure(
            text=self.t()[view] if view in self.t() else self.t()["dashboard"]
        )
        if view == "dashboard":
            self.render_dashboard()
            return
        student = StudentService.get_by_user_id(self.app.conn, self.user["id"])
        if student is None:
            self._empty_state(self.t()["no_student_profile"])
            return
        if view == "profile":
            self.render_profile(student)
        elif view == "grades":
            rows = GradeService.list_for_student(self.app.conn, student["id"])
            self.render_table(rows, ["subject_name", "semester_name", "current_score", "examination_score", "final_score", "note"])
        elif view == "schedule":
            rows = ScheduleService.list_for_group(self.app.conn, student["group_id"])
            self.render_table(rows, ["weekday", "starts_at", "ends_at", "subject_name", "teacher", "classroom"])
        elif view == "attendance":
            rows = AttendanceService.list_for_student(self.app.conn, student["id"])
            percentage = AttendanceService.percentage(rows)
            if percentage is not None:
                ctk.CTkLabel(
                    self.content,
                    text=f"{self.t()['attendance_rate']}: {percentage}%",
                    font=ctk.CTkFont(size=16, weight="bold"),
                ).grid(row=0, column=0, sticky="w", pady=(0, 8))
            self.render_table(
                rows,
                ["lesson_date", "subject_name", "status", "notes"],
                row=1 if percentage is not None else 0,
            )
        elif view == "payments":
            rows = PaymentService.list_for_student(self.app.conn, student["id"])
            self.render_table(rows, ["contract_number", "academic_year", "amount", "outstanding_balance", "contract_id"])
        elif view == "messages":
            rows = MessageService.list_for_student(self.app.conn, self.user["id"], student)
            self.render_table(rows, ["title", "published_at", "read", "body"])
            unread = [item for item in rows if not item["read"]]
            for index, row in enumerate(unread, start=1):
                if not row["read"]:
                    ctk.CTkButton(
                        self.content,
                        text=self.t()["mark_read"],
                        command=lambda mid=row["id"]: self.mark_read(mid),
                        fg_color="#176B87",
                    ).grid(row=index, column=0, sticky="e", pady=4)

    def render_dashboard(self) -> None:
        student = StudentService.get_by_user_id(self.app.conn, self.user["id"])
        name = student["full_name"] if student else self.user["username"]
        self.header_user.configure(text=f"{name}  ·  {self.t()['role_student']}")
        hero = ctk.CTkFrame(self.content, fg_color="#173B5D", corner_radius=18)
        hero.grid(row=0, column=0, padx=2, pady=(0, 18), sticky="ew")
        hero.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(
            hero,
            text=f"{self.t()['student_greeting']},\n{name}",
            font=ctk.CTkFont(size=27, weight="bold"),
            text_color="#FFFFFF",
            justify="left",
        ).grid(row=0, column=0, padx=26, pady=(24, 6), sticky="w")
        ctk.CTkLabel(
            hero,
            text=self.t()["student_welcome"],
            font=ctk.CTkFont(size=14),
            text_color="#D0DEEB",
        ).grid(row=1, column=0, padx=26, pady=(0, 22), sticky="w")
        banner_image = tk.PhotoImage(file=str(APP_IMAGE_PATH)).subsample(2, 2)
        self._banner_image = banner_image
        tk.Label(hero, image=banner_image, bg="#173B5D", bd=0).grid(
            row=0, column=1, rowspan=2, padx=20, pady=12, sticky="e"
        )
        if student is None:
            self._empty_state(self.t()["no_student_profile"], row=1)
            return

        grades = GradeService.list_for_student(self.app.conn, student["id"])
        average = (
            round(
                sum(float(row["final_score"]) for row in grades if row["final_score"] is not None)
                / sum(1 for row in grades if row["final_score"] is not None),
                1,
            )
            if any(row["final_score"] is not None for row in grades)
            else "—"
        )
        attendance = AttendanceService.list_for_student(self.app.conn, student["id"])
        attendance_rate = AttendanceService.percentage(attendance)
        schedule = ScheduleService.list_for_group(self.app.conn, student["group_id"])
        contracts = PaymentService.list_for_student(self.app.conn, student["id"])
        balance = sum(float(item["outstanding_balance"]) for item in contracts)
        cards = [
            ("avg_grade", average, "✧", "#176B87", "grades"),
            ("attendance_rate", f"{attendance_rate}%" if attendance_rate is not None else "—", "✓", "#287D65", "attendance"),
            ("schedule_count", len(schedule), "▦", "#7354A2", "schedule"),
            ("balance_due", f"{balance:,.2f} TJS", "¤", "#BD7B2A", "payments"),
        ]
        cards_frame = ctk.CTkFrame(self.content, fg_color="transparent")
        cards_frame.grid(row=1, column=0, sticky="ew", pady=(0, 18))
        for column in range(len(cards)):
            cards_frame.grid_columnconfigure(column, weight=1, uniform="student-card")
        for column, (key, value, icon, accent, target) in enumerate(cards):
            self._metric_card(cards_frame, column, key, value, icon, accent, target)
        ctk.CTkLabel(
            self.content,
            text=self.t()["quick_actions"],
            font=ctk.CTkFont(size=18, weight="bold"),
            text_color="#172B4D",
        ).grid(row=2, column=0, padx=4, pady=(6, 10), sticky="w")
        actions = ctk.CTkFrame(self.content, fg_color="transparent")
        actions.grid(row=3, column=0, sticky="ew")
        for column, key in enumerate(("grades", "schedule", "messages")):
            actions.grid_columnconfigure(column, weight=1, uniform="student-actions")
            ctk.CTkButton(
                actions,
                text=f"{self.t()[key]}  →",
                command=lambda target=key: self.show_view(target),
                height=48,
                corner_radius=11,
                fg_color="#FFFFFF",
                hover_color="#E9F0F6",
                text_color="#173B5D",
                border_width=1,
                border_color="#E0E8F0",
                anchor="w",
            ).grid(row=0, column=column, padx=(0, 10), sticky="ew")

    def _metric_card(
        self, parent: ctk.CTkFrame, column: int, key: str, value: Any,
        icon: str, accent: str, target: str
    ) -> None:
        card = ctk.CTkFrame(parent, fg_color="#FFFFFF", corner_radius=14)
        card.grid(row=0, column=column, padx=(0, 10), sticky="nsew")
        ctk.CTkLabel(
            card, text=icon, text_color=accent, font=ctk.CTkFont(size=22, weight="bold")
        ).pack(anchor="w", padx=16, pady=(14, 4))
        ctk.CTkLabel(
            card, text=self.t()[key], text_color="#718198", font=ctk.CTkFont(size=12)
        ).pack(anchor="w", padx=16)
        ctk.CTkButton(
            card,
            text=str(value),
            command=lambda: self.show_view(target),
            fg_color="transparent",
            hover_color="#F1F5F9",
            text_color="#172B4D",
            font=ctk.CTkFont(size=22, weight="bold"),
            anchor="w",
            height=42,
        ).pack(fill="x", padx=8, pady=(1, 10))

    def _empty_state(self, text: str, row: int = 0) -> None:
        empty = ctk.CTkFrame(self.content, fg_color="#FFFFFF", corner_radius=14)
        empty.grid(row=row, column=0, padx=2, pady=4, sticky="ew")
        ctk.CTkLabel(
            empty, text="◌", font=ctk.CTkFont(size=34), text_color="#8AA0B8"
        ).pack(pady=(24, 5))
        ctk.CTkLabel(
            empty,
            text=text,
            font=ctk.CTkFont(size=14),
            text_color="#63748A",
            wraplength=620,
        ).pack(padx=24, pady=(0, 24))

    def render_profile(self, student: dict[str, Any]) -> None:
        card = ctk.CTkFrame(self.content, fg_color="#FFFFFF", corner_radius=16)
        card.grid(row=0, column=0, sticky="ew", padx=2, pady=3)
        details = [
            ("full_name", student.get("full_name")),
            ("student_number", student.get("student_number")),
            ("group_name", student.get("group_name") or student.get("group_id")),
            ("specialty_name", student.get("specialty_name") or "—"),
            ("course_name", student.get("course_name") or "—"),
            ("email", student.get("email") or "—"),
            ("phone", student.get("phone") or "—"),
            ("birth_date", student.get("birth_date") or "—"),
        ]
        for row, (key, value) in enumerate(details):
            ctk.CTkLabel(
                card, text=self.t().get(key, key), text_color="#718198",
                font=ctk.CTkFont(size=12),
            ).grid(row=row, column=0, padx=20, pady=10, sticky="w")
            ctk.CTkLabel(
                card, text=str(value), text_color="#172B4D",
                font=ctk.CTkFont(size=14, weight="bold"),
            ).grid(row=row, column=1, padx=20, pady=10, sticky="w")

    def render_table(
        self, rows: list[dict[str, Any]], fields: list[str], row: int = 0
    ) -> None:
        self.current_fields = fields
        if not rows:
            self._empty_state(self.t()["no_data"], row=row)
            return
        table_frame = ctk.CTkFrame(self.content, fg_color="#FFFFFF", corner_radius=14)
        table_frame.grid(row=row, column=0, sticky="nsew", padx=2, pady=4)
        table_frame.grid_columnconfigure(0, weight=1)
        table_frame.grid_rowconfigure(0, weight=1)
        style = ttk.Style()
        style.configure(
            "Cabinet.Treeview",
            background="#FFFFFF",
            foreground="#263A52",
            fieldbackground="#FFFFFF",
            rowheight=34,
            font=("Segoe UI", 10),
        )
        style.configure(
            "Cabinet.Treeview.Heading",
            background="#EAF0F6",
            foreground="#173B5D",
            font=("Segoe UI", 10, "bold"),
            relief="flat",
        )
        style.map("Cabinet.Treeview", background=[("selected", "#D8EAF0")])
        table = ttk.Treeview(table_frame, columns=fields, show="headings", style="Cabinet.Treeview")
        self.data_table = table
        yscroll = ttk.Scrollbar(table_frame, orient="vertical", command=table.yview)
        xscroll = ttk.Scrollbar(table_frame, orient="horizontal", command=table.xview)
        table.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        table.grid(row=0, column=0, sticky="nsew", padx=(12, 0), pady=(12, 0))
        yscroll.grid(row=0, column=1, sticky="ns", padx=(0, 12), pady=(12, 0))
        xscroll.grid(row=1, column=0, sticky="ew", padx=(12, 0), pady=(0, 12))
        for field in fields:
            table.heading(field, text=self.t().get(field, field.replace("_", " ").title()))
            table.column(field, width=150, anchor="w", stretch=True)
        for row in rows:
            values = [row.get(field) for field in fields]
            values = [
                self._display_value(field, value) for field, value in zip(fields, values)
            ]
            table.insert("", "end", values=[str(v) if v is not None else "" for v in values])

    def _display_value(self, field: str, value: Any) -> Any:
        if field == "status":
            return self.t().get(f"attendance_{value}", value)
        if field in {"read", "is_published", "is_active"}:
            return self.t()["yes"] if value else self.t()["no"]
        if field == "weekday" and isinstance(value, int):
            weekdays = (
                ("Понедельник", "Вторник", "Среда", "Четверг", "Пятница", "Суббота", "Воскресенье")
                if self.language == "ru"
                else ("Душанбе", "Сешанбе", "Чоршанбе", "Панҷшанбе", "Ҷумъа", "Шанбе", "Якшанбе")
            )
            return weekdays[value] if value in range(7) else value
        return value

    def _run_form(self, action: Any) -> None:
        try:
            action()
        except (ValueError, TypeError, sqlite3.Error) as exc:
            self.app.conn.rollback()
            logger.exception("Form submission failed")
            detail = str(exc)
            known_errors = {
                "Payment exceeds remaining balance.": self.t()["payment_exceeds_balance"],
                "Please enter final score or configure official grade weights.": self.t()["grade_weights_missing"],
            }
            localized_detail = (
                known_errors.get(detail, detail)
                if detail in self.t().values()
                else known_errors.get(detail, self.t()["invalid_value"])
            )
            messagebox.showerror(
                self.t()["error"],
                localized_detail,
                parent=self,
            )

    def _defer_dialog_close(
        self, dialog: ctk.CTkToplevel, refresh: Any
    ) -> None:
        def finish() -> None:
            if dialog.winfo_exists():
                dialog.destroy()
            refresh()

        dialog.withdraw()
        dialog.after_idle(finish)

    def mark_read(self, message_id: int) -> None:
        MessageService.mark_read(self.app.conn, message_id, self.user["id"])
        self.show_view("messages")

    def change_password(self) -> None:
        dlg = ctk.CTkToplevel(self)
        dlg.title(self.t()["new_password"])
        dlg.geometry("360x220")
        dlg.transient(self)
        old = ctk.CTkEntry(dlg, placeholder_text=self.t()["current_password"], show="*")
        old.pack(fill="x", padx=20, pady=(18, 8))
        new = ctk.CTkEntry(dlg, placeholder_text=self.t()["new_password"], show="*")
        new.pack(fill="x", padx=20, pady=8)

        def save() -> None:
            user = self.app.conn.execute("SELECT * FROM users WHERE id = ?", (self.user["id"],)).fetchone()
            if user is None or not AuthenticationService.verify_password(user["password_hash"], old.get()):
                messagebox.showerror(
                    self.t()["error"], self.t()["current_password_error"], parent=dlg
                )
                return
            if len(new.get()) < 8:
                messagebox.showerror(
                    self.t()["error"], self.t()["password_min_length"], parent=dlg
                )
                return
            self.app.conn.execute(
                "UPDATE users SET password_hash = ? WHERE id = ?",
                (AuthenticationService.hash_password(new.get()), self.user["id"]),
            )
            self.app.conn.commit()
            messagebox.showinfo(
                self.t()["success"], self.t()["password_changed"], parent=dlg
            )
            self._defer_dialog_close(dlg, lambda: None)
        ctk.CTkButton(
            dlg, text=self.t()["save"], command=lambda: self._run_form(save),
            fg_color="#176B87",
        ).pack(pady=18)


class AdminDashboard(StudentDashboard):
    """Administrative dashboard with data management."""

    def build_sidebar(self) -> None:
        sidebar = ctk.CTkFrame(self, width=246, corner_radius=0, fg_color="#102A4C")
        sidebar.grid(row=0, column=0, sticky="nsew")
        sidebar.grid_propagate(False)
        sidebar.grid_rowconfigure(10, weight=1)
        ctk.CTkLabel(
            sidebar,
            text="✦  " + self.t()["app_name"],
            text_color="#FFFFFF",
            font=ctk.CTkFont(size=19, weight="bold"),
        ).grid(row=0, column=0, padx=20, pady=(26, 34), sticky="w")
        options = [
            ("dashboard", "⌂", self.t()["dashboard"]),
            ("students", "◉", self.t()["students"]),
            ("catalogs", "▤", self.t()["catalogs"]),
            ("grades", "✧", self.t()["grades"]),
            ("schedule", "▦", self.t()["schedule"]),
            ("attendance", "✓", self.t()["attendance"]),
            ("payments", "¤", self.t()["payments"]),
            ("messages", "✉", self.t()["messages"]),
            ("profile", "◉", self.t()["profile"]),
        ]
        self.nav_buttons = {}
        for row, (key, icon, label) in enumerate(options, start=1):
            button = ctk.CTkButton(
                sidebar,
                text=f"{icon}    {label}",
                command=lambda value=key: self.show_view(value),
                fg_color="#1D4268" if key == "dashboard" else "transparent",
                hover_color="#244C78",
                text_color="#E9F0F8",
                anchor="w",
                height=44,
                corner_radius=10,
                font=ctk.CTkFont(size=14),
            )
            button.grid(row=row, column=0, padx=12, pady=3, sticky="ew")
            self.nav_buttons[key] = button
        ctk.CTkLabel(
            sidebar,
            text=f"{self.user['username']}  ·  {self.t()['role_admin']}",
            text_color="#B9C9DC",
            font=ctk.CTkFont(size=12),
        ).grid(row=11, column=0, padx=20, pady=(12, 10), sticky="w")
        ctk.CTkButton(
            sidebar,
            text=self.t()["password_changed"],
            command=self.change_password,
            fg_color="#1D4268",
            hover_color="#244C78",
            height=40,
            corner_radius=10,
        ).grid(row=12, column=0, padx=12, pady=4, sticky="ew")
        ctk.CTkButton(
            sidebar,
            text="↪   " + self.t()["logout"],
            command=self.app.logout,
            fg_color="#A64045",
            hover_color="#8F3339",
            height=42,
            corner_radius=10,
        ).grid(row=13, column=0, padx=12, pady=(4, 18), sticky="ew")

    def show_view(self, view: str) -> None:
        for child in self.content.winfo_children():
            child.destroy()
        for key, button in self.nav_buttons.items():
            button.configure(
                fg_color="#1D4268" if key == view else "transparent"
            )
        self.header_user.configure(
            text=f"{self.user['username']}  ·  {self.t()['role_admin']}"
        )
        self.page_title.configure(
            text=self.t()["admin_profile_title"]
            if view == "profile"
            else self.t().get(view, self.t()["dashboard"])
        )
        if view == "dashboard":
            self.render_dashboard()
            return
        if view == "profile":
            self.render_admin_profile()
            return
        if view == "catalogs":
            self.render_catalogs()
            return
        toolbar = ctk.CTkFrame(self.content, fg_color="transparent")
        toolbar.grid(row=0, column=0, sticky="ew", pady=(0, 12))
        toolbar.grid_columnconfigure(0, weight=1)
        query = ctk.CTkEntry(toolbar, placeholder_text=self.t()["search"])
        query.grid(row=0, column=0, padx=(0, 8), sticky="ew")
        ctk.CTkButton(
            toolbar,
            text=self.t()["search"],
            command=lambda: self.load_search(view, query.get()),
            fg_color="#176B87",
            width=110,
        ).grid(row=0, column=1, padx=(0, 8))
        ctk.CTkButton(
            toolbar,
            text="+  " + self.t()["add"],
            command=lambda: self.open_add_form(view),
            fg_color="#287D65",
            width=120,
        ).grid(row=0, column=2, padx=(0, 8))
        if view == "students":
            ctk.CTkButton(
                toolbar,
                text=self.t()["edit_student"],
                command=self.edit_selected_student,
                fg_color="#FFFFFF",
                text_color="#173B5D",
                hover_color="#E9F0F6",
                border_width=1,
                border_color="#DCE5EE",
                width=150,
            ).grid(row=0, column=3, padx=(0, 8))
            ctk.CTkButton(
                toolbar,
                text=self.t()["toggle_student"],
                command=self.toggle_selected_student,
                fg_color="#FFFFFF",
                text_color="#173B5D",
                hover_color="#E9F0F6",
                border_width=1,
                border_color="#DCE5EE",
                width=145,
            ).grid(row=0, column=4, padx=(0, 8))
            export_column = 5
        else:
            export_column = 3
        ctk.CTkButton(
            toolbar,
            text=self.t()["export"],
            command=lambda: self.export_view(view),
            fg_color="#FFFFFF",
            text_color="#173B5D",
            hover_color="#E9F0F6",
            border_width=1,
            border_color="#DCE5EE",
            width=145,
        ).grid(row=0, column=export_column)
        self.load_search(view, query.get())

    def render_admin_profile(self) -> None:
        account = self.app.conn.execute(
            "SELECT id, username, role, is_active, created_at "
            "FROM users WHERE id = ?",
            (self.user["id"],),
        ).fetchone()
        if account is None:
            raise RuntimeError("The signed-in administrator account no longer exists.")

        hero = ctk.CTkFrame(self.content, fg_color="#173B5D", corner_radius=18)
        hero.grid(row=0, column=0, padx=2, pady=(0, 16), sticky="ew")
        hero.grid_columnconfigure(1, weight=1)
        initials = account["username"][:1].upper()
        ctk.CTkLabel(
            hero,
            text=initials,
            width=66,
            height=66,
            corner_radius=33,
            fg_color="#287D65",
            text_color="#FFFFFF",
            font=ctk.CTkFont(size=27, weight="bold"),
        ).grid(row=0, column=0, rowspan=2, padx=(22, 16), pady=20)
        ctk.CTkLabel(
            hero,
            text=account["username"],
            text_color="#FFFFFF",
            font=ctk.CTkFont(size=24, weight="bold"),
        ).grid(row=0, column=1, padx=(0, 20), pady=(22, 2), sticky="sw")
        ctk.CTkLabel(
            hero,
            text=self.t()["role_admin"],
            text_color="#D0DEEB",
            font=ctk.CTkFont(size=14),
        ).grid(row=1, column=1, padx=(0, 20), pady=(0, 22), sticky="nw")

        details = [
            ("account_id", account["id"]),
            ("username", account["username"]),
            ("role_admin", self.t()["role_admin"]),
            (
                "account_status",
                self.t()["active"] if account["is_active"] else self.t()["inactive"],
            ),
            ("created_at", account["created_at"]),
        ]
        card = ctk.CTkFrame(self.content, fg_color="#FFFFFF", corner_radius=14)
        card.grid(row=1, column=0, padx=2, sticky="ew")
        card.grid_columnconfigure(1, weight=1)
        for row, (key, value) in enumerate(details):
            ctk.CTkLabel(
                card,
                text=self.t()[key],
                text_color="#718198",
                font=ctk.CTkFont(size=13),
            ).grid(row=row, column=0, padx=(20, 26), pady=12, sticky="w")
            ctk.CTkLabel(
                card,
                text=str(value),
                text_color="#172B4D",
                font=ctk.CTkFont(size=14, weight="bold"),
            ).grid(row=row, column=1, padx=20, pady=12, sticky="w")

    def load_search(self, view: str, term: str) -> None:
        rows: list[dict[str, Any]] = []
        if view == "students":
            rows = StudentService.list(self.app.conn, term)
            self.render_table(
                rows,
                ["id", "full_name", "student_number", "username", "group_name", "email", "is_active"],
                row=1,
            )
        elif view == "grades":
            rows = GradeService.list_all(self.app.conn)
            if term:
                rows = [r for r in rows if term.lower() in (r.get("full_name") or "").lower() or term.lower() in (r.get("subject_name") or "").lower()]
            self.render_table(rows, ["full_name", "student_number", "subject_name", "semester_name", "final_score", "note"], row=1)
        elif view == "schedule":
            rows = ScheduleService.list_all(self.app.conn)
            if term:
                rows = [r for r in rows if term.lower() in (r.get("group_name") or "").lower() or term.lower() in (r.get("subject_name") or "").lower()]
            self.render_table(rows, ["group_name", "subject_name", "weekday", "starts_at", "ends_at", "teacher", "classroom"], row=1)
        elif view == "attendance":
            rows = self.app.conn.execute("SELECT a.*, s.name AS subject_name, st.full_name, st.student_number FROM attendance a JOIN subjects s ON s.id = a.subject_id JOIN students st ON st.id = a.student_id ORDER BY a.lesson_date DESC").fetchall()
            if term:
                rows = [r for r in rows if term.lower() in (r["full_name"] or "").lower() or term.lower() in (r["subject_name"] or "").lower()]
            rows = [dict(r) for r in rows]
            self.render_table(rows, ["full_name", "student_number", "subject_name", "lesson_date", "status", "notes"], row=1)
        elif view == "payments":
            rows = self.app.conn.execute("SELECT p.*, c.contract_number, s.full_name FROM payments p JOIN contracts c ON c.id = p.contract_id JOIN students s ON s.id = c.student_id ORDER BY p.payment_date DESC").fetchall()
            if term:
                rows = [r for r in rows if term.lower() in (r["full_name"] or "").lower() or term.lower() in (r["contract_number"] or "").lower()]
            rows = [dict(r) for r in rows]
            self.render_table(rows, ["full_name", "contract_number", "amount", "payment_date", "receipt_number"], row=1)
        elif view == "messages":
            rows = MessageService.list_all(self.app.conn)
            if term:
                rows = [r for r in rows if term.lower() in (r["title"] or "").lower() or term.lower() in (r["body"] or "").lower()]
            self.render_table(rows, ["title", "group_name", "published_at", "is_published", "body"], row=1)
        self.current_rows = rows

    def _selected_student_id(self) -> int | None:
        if self.data_table is None:
            messagebox.showerror(
                self.t()["error"], self.t()["student_selected"], parent=self
            )
            return None
        selection = self.data_table.selection()
        if not selection:
            messagebox.showerror(
                self.t()["error"], self.t()["student_selected"], parent=self
            )
            return None
        return int(self.data_table.item(selection[0], "values")[0])

    def toggle_selected_student(self) -> None:
        student_id = self._selected_student_id()
        if student_id is None:
            return
        row = self.app.conn.execute(
            "SELECT u.is_active FROM students s JOIN users u ON u.id = s.user_id "
            "WHERE s.id = ?",
            (student_id,),
        ).fetchone()
        if row is None:
            messagebox.showerror(
                self.t()["error"], self.t()["invalid_value"], parent=self
            )
            return
        if row["is_active"]:
            StudentService.deactivate_student(self.app.conn, student_id)
            result = self.t()["student_deactivated"]
        else:
            StudentService.activate_student(self.app.conn, student_id)
            result = self.t()["student_activated"]
        self.load_search("students", "")
        messagebox.showinfo(self.t()["success"], result, parent=self)

    def edit_selected_student(self) -> None:
        student_id = self._selected_student_id()
        if student_id is None:
            return
        student = next(
            (row for row in StudentService.list(self.app.conn) if row["id"] == student_id),
            None,
        )
        if student is None:
            messagebox.showerror(
                self.t()["error"], self.t()["invalid_value"], parent=self
            )
            return
        dlg = ctk.CTkToplevel(self)
        dlg.title(self.t()["edit_student"])
        dlg.geometry("460x470")
        dlg.transient(self)
        dlg.grid_columnconfigure(0, weight=1)
        fields = [
            "student_number", "full_name", "group_id", "specialty_id",
            "course_id", "email", "phone", "birth_date",
        ]
        references = {
            "group_id": ("groups", False),
            "specialty_id": ("specialties", True),
            "course_id": ("courses", True),
        }
        entries: dict[str, Any] = {}
        for index, key in enumerate(fields):
            if key in references:
                table, optional = references[key]
                self._add_reference_field(dlg, key, index, entries, table, optional)
                value = student.get(key)
                if value is not None:
                    option = next(
                        item for item in entries[key].cget("values")
                        if item.startswith(f"{value} · ")
                    )
                    entries[key].set(option)
                continue
            entry = ctk.CTkEntry(
                dlg,
                height=40,
                placeholder_text=self.t().get(key, key.replace("_", " ").title()),
            )
            if student.get(key) is not None:
                entry.insert(0, str(student[key]))
            entry.grid(row=index, column=0, padx=18, pady=6, sticky="ew")
            entries[key] = entry

        def save() -> None:
            payload = {key: widget.get().strip() for key, widget in entries.items()}
            if not payload["student_number"] or not payload["full_name"]:
                raise ValueError(self.t()["required_fields"])
            values = {
                "student_number": payload["student_number"],
                "full_name": payload["full_name"],
                "group_id": self._reference_id(entries["group_id"]),
                "specialty_id": self._reference_id(entries["specialty_id"], optional=True),
                "course_id": self._reference_id(entries["course_id"], optional=True),
                "email": payload["email"] or None,
                "phone": payload["phone"] or None,
                "birth_date": payload["birth_date"] or None,
            }
            StudentService.update_student(self.app.conn, student_id, values)
            self._defer_dialog_close(
                dlg, lambda: self.load_search("students", "")
            )

        ctk.CTkButton(
            dlg,
            text=self.t()["save"],
            command=lambda: self._run_form(save),
            fg_color="#176B87",
        ).grid(row=len(fields), column=0, padx=18, pady=14, sticky="ew")

    def render_dashboard(self) -> None:
        conn = self.app.conn
        counts = {
            "total_students": conn.execute("SELECT COUNT(*) FROM students").fetchone()[0],
            "total_groups": conn.execute("SELECT COUNT(*) FROM groups").fetchone()[0],
            "total_subjects": conn.execute("SELECT COUNT(*) FROM subjects").fetchone()[0],
        }
        balance = conn.execute(
            "SELECT COALESCE(SUM(c.amount - COALESCE(p.paid, 0)), 0) "
            "FROM contracts c LEFT JOIN ("
            "SELECT contract_id, SUM(amount) AS paid FROM payments GROUP BY contract_id"
            ") p ON p.contract_id = c.id"
        ).fetchone()[0]
        hero = ctk.CTkFrame(self.content, fg_color="#173B5D", corner_radius=18)
        hero.grid(row=0, column=0, padx=2, pady=(0, 18), sticky="ew")
        hero.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(
            hero,
            text=f"{self.t()['admin_greeting']},\n{self.user['username']}",
            font=ctk.CTkFont(size=27, weight="bold"),
            text_color="#FFFFFF",
            justify="left",
        ).grid(row=0, column=0, padx=26, pady=(24, 6), sticky="w")
        ctk.CTkLabel(
            hero,
            text=self.t()["admin_welcome"],
            font=ctk.CTkFont(size=14),
            text_color="#D0DEEB",
        ).grid(row=1, column=0, padx=26, pady=(0, 22), sticky="w")
        banner_image = tk.PhotoImage(file=str(APP_IMAGE_PATH)).subsample(2, 2)
        self._banner_image = banner_image
        tk.Label(hero, image=banner_image, bg="#173B5D", bd=0).grid(
            row=0, column=1, rowspan=2, padx=20, pady=12, sticky="e"
        )
        cards_frame = ctk.CTkFrame(self.content, fg_color="transparent")
        cards_frame.grid(row=1, column=0, sticky="ew", pady=(0, 18))
        cards = [
            ("total_students", counts["total_students"], "◉", "#176B87", "students"),
            ("total_groups", counts["total_groups"], "▤", "#7354A2", "catalogs"),
            ("total_subjects", counts["total_subjects"], "✧", "#287D65", "catalogs"),
            ("balance_due", f"{float(balance):,.2f} TJS", "¤", "#BD7B2A", "payments"),
        ]
        for column in range(len(cards)):
            cards_frame.grid_columnconfigure(column, weight=1, uniform="admin-card")
        for column, (key, value, icon, accent, target) in enumerate(cards):
            self._metric_card(cards_frame, column, key, value, icon, accent, target)
        ctk.CTkLabel(
            self.content,
            text=self.t()["quick_actions"],
            font=ctk.CTkFont(size=18, weight="bold"),
            text_color="#172B4D",
        ).grid(row=2, column=0, padx=4, pady=(6, 10), sticky="w")
        actions = ctk.CTkFrame(self.content, fg_color="transparent")
        actions.grid(row=3, column=0, sticky="ew")
        actions_to_run = [
            ("add_student", self.add_or_edit_student),
            ("add_grade", self.add_or_edit_grade),
            ("add_schedule", self.add_or_edit_schedule),
            ("catalogs", lambda: self.show_view("catalogs")),
        ]
        for column, (key, callback) in enumerate(actions_to_run):
            actions.grid_columnconfigure(column, weight=1, uniform="admin-actions")
            ctk.CTkButton(
                actions,
                text=f"{self.t()[key]}  →",
                command=callback,
                height=48,
                corner_radius=11,
                fg_color="#FFFFFF",
                hover_color="#E9F0F6",
                text_color="#173B5D",
                border_width=1,
                border_color="#E0E8F0",
                anchor="w",
            ).grid(row=0, column=column, padx=(0, 10), sticky="ew")
        latest = MessageService.list_all(conn)[:3]
        ctk.CTkLabel(
            self.content,
            text=self.t()["recent_messages"],
            font=ctk.CTkFont(size=18, weight="bold"),
            text_color="#172B4D",
        ).grid(row=4, column=0, padx=4, pady=(22, 10), sticky="w")
        if latest:
            for row, item in enumerate(latest, start=5):
                card = ctk.CTkFrame(self.content, fg_color="#FFFFFF", corner_radius=12)
                card.grid(row=row, column=0, sticky="ew", pady=4)
                ctk.CTkLabel(
                    card, text=item["title"], text_color="#172B4D",
                    font=ctk.CTkFont(size=14, weight="bold"),
                ).pack(anchor="w", padx=16, pady=(12, 3))
                ctk.CTkLabel(
                    card,
                    text=f"{item.get('group_name') or self.t()['students']}  ·  {item['published_at']}",
                    text_color="#718198",
                    font=ctk.CTkFont(size=12),
                ).pack(anchor="w", padx=16, pady=(0, 12))
        else:
            self._empty_state(self.t()["no_data"], row=5)

    def render_catalogs(self) -> None:
        self._catalogs = {
            "groups": ["id", "name"],
            "subjects": ["id", "name", "code"],
            "specialties": ["id", "name"],
            "courses": ["id", "name", "year_number"],
            "semesters": ["id", "name", "academic_year", "term", "starts_on", "ends_on"],
            "contracts": [
                "id", "contract_number", "full_name", "amount", "currency",
                "academic_year", "outstanding_balance",
            ],
        }
        self._catalog_labels = {
            self.t()["catalog_groups"]: "groups",
            self.t()["catalog_subjects"]: "subjects",
            self.t()["catalog_specialties"]: "specialties",
            self.t()["catalog_courses"]: "courses",
            self.t()["catalog_semesters"]: "semesters",
            self.t()["catalog_contracts"]: "contracts",
        }
        toolbar = ctk.CTkFrame(self.content, fg_color="transparent")
        toolbar.grid(row=0, column=0, sticky="ew", pady=(0, 10))
        toolbar.grid_columnconfigure(0, weight=1)
        self.catalog_choice = ctk.StringVar(value=self.t()["catalog_groups"])
        ctk.CTkOptionMenu(
            toolbar,
            variable=self.catalog_choice,
            values=list(self._catalog_labels),
            command=lambda _: self.load_catalogs(),
            fg_color="#176B87",
        ).grid(row=0, column=0, sticky="w")
        ctk.CTkButton(
            toolbar,
            text="+  " + self.t()["add"],
            command=self.add_catalog_item,
            fg_color="#287D65",
        ).grid(row=0, column=1, padx=8)
        ctk.CTkLabel(
            self.content,
            text=self.t()["catalogs_hint"],
            text_color="#718198",
            wraplength=850,
            justify="left",
        ).grid(row=1, column=0, sticky="w", pady=(0, 10))
        self.load_catalogs()

    def load_catalogs(self) -> None:
        table = self._catalog_labels[self.catalog_choice.get()]
        fields = self._catalogs[table]
        if table == "contracts":
            rows = [
                dict(row)
                for row in self.app.conn.execute(
                    "SELECT c.id, c.contract_number, s.full_name, c.amount, "
                    "c.currency, c.academic_year, "
                    "c.amount - COALESCE(SUM(p.amount), 0) AS outstanding_balance "
                    "FROM contracts c JOIN students s ON s.id = c.student_id "
                    "LEFT JOIN payments p ON p.contract_id = c.id "
                    "GROUP BY c.id ORDER BY c.academic_year DESC, s.full_name"
                ).fetchall()
            ]
        else:
            rows = [
                dict(row)
                for row in self.app.conn.execute(
                    f"SELECT {', '.join(fields)} FROM {table} ORDER BY id"
                ).fetchall()
            ]
        self.current_rows = rows
        self.render_table(rows, fields, row=2)

    def add_catalog_item(self) -> None:
        table = self._catalog_labels[self.catalog_choice.get()]
        if table == "contracts":
            self.add_contract()
            return
        fields = {
            "groups": ["name"],
            "subjects": ["name", "code"],
            "specialties": ["name"],
            "courses": ["name", "year_number"],
            "semesters": ["name", "academic_year", "term", "starts_on", "ends_on"],
        }[table]
        dlg = ctk.CTkToplevel(self)
        dlg.title(self.catalog_choice.get())
        dlg.geometry("480x430" if table == "semesters" else "440x280")
        dlg.transient(self)
        dlg.grid_columnconfigure(0, weight=1)
        entries: dict[str, ctk.CTkEntry] = {}
        for row, field in enumerate(fields):
            entry = ctk.CTkEntry(
                dlg,
                height=42,
                placeholder_text=self.t().get(field, field.replace("_", " ").title()),
            )
            entry.grid(row=row, column=0, padx=22, pady=7, sticky="ew")
            entries[field] = entry

        def save() -> None:
            values = {field: entry.get().strip() for field, entry in entries.items()}
            if any(not values[field] for field in fields):
                raise ValueError(self.t()["required_fields"])
            if table == "courses":
                values["year_number"] = int(values["year_number"])
            elif table == "semesters":
                values["term"] = int(values["term"])
                if values["term"] not in (1, 2):
                    raise ValueError(self.t()["invalid_value"])
                date.fromisoformat(values["starts_on"])
                date.fromisoformat(values["ends_on"])
                if values["starts_on"] > values["ends_on"]:
                    raise ValueError(self.t()["invalid_value"])
            columns = ", ".join(fields)
            placeholders = ", ".join("?" for _ in fields)
            self.app.conn.execute(
                f"INSERT INTO {table} ({columns}) VALUES ({placeholders})",
                tuple(values[field] for field in fields),
            )
            self.app.conn.commit()
            self._defer_dialog_close(dlg, self.load_catalogs)

        ctk.CTkButton(
            dlg, text=self.t()["save"], command=lambda: self._run_form(save),
            fg_color="#176B87",
        ).grid(row=len(fields), column=0, padx=22, pady=16, sticky="ew")

    def add_contract(self) -> None:
        students = StudentService.list(self.app.conn)
        if not students:
            messagebox.showerror(
                self.t()["error"], self.t()["no_data"], parent=self
            )
            return
        dlg = ctk.CTkToplevel(self)
        dlg.title(self.t()["catalog_contracts"])
        dlg.geometry("480x380")
        dlg.transient(self)
        dlg.grid_columnconfigure(0, weight=1)
        student_options = [
            f"{row['id']} · {row['full_name']} ({row['student_number']})"
            for row in students
        ]
        ctk.CTkLabel(
            dlg, text=self.t()["student_id"], text_color="#52647C"
        ).grid(row=0, column=0, padx=22, pady=(20, 4), sticky="w")
        selected_student = ctk.StringVar(value=student_options[0])
        ctk.CTkOptionMenu(
            dlg, variable=selected_student, values=student_options
        ).grid(row=1, column=0, padx=22, pady=(0, 10), sticky="ew")
        fields = ["amount", "currency", "contract_number", "academic_year"]
        entries: dict[str, ctk.CTkEntry] = {}
        for row, key in enumerate(fields, start=2):
            entry = ctk.CTkEntry(
                dlg, placeholder_text=self.t().get(key, key.replace("_", " ").title())
            )
            if key == "currency":
                entry.insert(0, "TJS")
            entry.grid(row=row, column=0, padx=22, pady=6, sticky="ew")
            entries[key] = entry

        def save() -> None:
            values = {key: entry.get().strip() for key, entry in entries.items()}
            if any(not values[key] for key in fields):
                raise ValueError(self.t()["required_fields"])
            amount = float(values["amount"])
            if amount <= 0:
                raise ValueError(self.t()["invalid_value"])
            student_id = int(selected_student.get().split(" · ", 1)[0])
            self.app.conn.execute(
                "INSERT INTO contracts(student_id, amount, currency, contract_number, academic_year) "
                "VALUES (?, ?, ?, ?, ?)",
                (
                    student_id, amount, values["currency"], values["contract_number"],
                    values["academic_year"],
                ),
            )
            self.app.conn.commit()
            def refresh() -> None:
                self.load_catalogs()
                messagebox.showinfo(
                    self.t()["success"], self.t()["contract_created"], parent=self
                )

            self._defer_dialog_close(dlg, refresh)

        ctk.CTkButton(
            dlg, text=self.t()["save"], command=lambda: self._run_form(save),
            fg_color="#176B87",
        ).grid(row=len(fields) + 2, column=0, padx=22, pady=14, sticky="ew")

    def export_view(self, view: str) -> None:
        target = filedialog.asksaveasfilename(
            parent=self,
            title=self.t()["export"],
            defaultextension=".xlsx",
            filetypes=[(self.t()["excel_files"], "*.xlsx")],
            initialfile=f"{self.t().get(view, self.t()['report'])}.xlsx",
        )
        if not target:
            return
        try:
            rows = getattr(self, "current_rows", [])
            fields = getattr(self, "current_fields", [])
            localized_rows = [
                {
                    self.t().get(key, key.replace("_", " ").title()):
                    self._display_value(key, record.get(key))
                    for key in fields
                }
                for record in rows
            ]
            ExcelExporter.export(
                localized_rows, target, self.t()["report"], self.t()["no_data"]
            )
            messagebox.showinfo(self.t()["success"], target, parent=self)
        except (OSError, ValueError):
            logger.exception("Excel export failed")
            messagebox.showerror(self.t()["error"], self.t()["invalid_value"], parent=self)

    def _add_reference_field(
        self,
        dialog: ctk.CTkToplevel,
        key: str,
        row: int,
        entries: dict[str, Any],
        table: str,
        optional: bool = False,
    ) -> None:
        queries = {
            "students": (
                "SELECT id, full_name || ' — ' || student_number AS label "
                "FROM students ORDER BY full_name"
            ),
            "groups": "SELECT id, name AS label FROM groups ORDER BY name",
            "subjects": "SELECT id, name AS label FROM subjects ORDER BY name",
            "specialties": "SELECT id, name AS label FROM specialties ORDER BY name",
            "courses": "SELECT id, name AS label FROM courses ORDER BY year_number",
            "semesters": (
                "SELECT id, name || ' — ' || academic_year || ' / ' || term AS label "
                "FROM semesters ORDER BY academic_year DESC, term"
            ),
            "contracts": (
                "SELECT id, contract_number || ' — ' || academic_year AS label "
                "FROM contracts ORDER BY academic_year DESC"
            ),
        }
        if table not in queries:
            raise ValueError("Unsupported reference type")
        options = [
            f"{record['id']} · {record['label']}"
            for record in self.app.conn.execute(queries[table]).fetchall()
        ]
        if optional:
            values = [self.t()["select_none"], *options]
            initial_value = self.t()["select_none"]
        elif options:
            values = options
            initial_value = options[0]
        else:
            values = [self.t()["no_data"]]
            initial_value = self.t()["no_data"]
        selection = ctk.StringVar(value=initial_value)
        widget = ctk.CTkOptionMenu(
            dialog,
            variable=selection,
            values=values,
            fg_color="#176B87",
            button_color="#12566D",
            height=40,
        )
        widget.grid(row=row, column=0, padx=18, pady=6, sticky="ew")
        if not optional and not options:
            widget.configure(state="disabled")
        entries[key] = widget

    def _reference_id(self, widget: Any, optional: bool = False) -> int | None:
        selected = widget.get()
        if selected == self.t()["select_none"] and optional:
            return None
        if " · " not in selected:
            raise ValueError(self.t()["required_fields"])
        return int(selected.split(" · ", 1)[0])

    def open_add_form(self, view: str) -> None:
        if view == "students":
            self.add_or_edit_student()
        elif view == "grades":
            self.add_or_edit_grade()
        elif view == "schedule":
            self.add_or_edit_schedule()
        elif view == "attendance":
            self.add_or_edit_attendance()
        elif view == "payments":
            self.add_or_edit_payment()
        elif view == "messages":
            self.add_or_edit_message()

    def add_or_edit_student(self) -> None:
        dlg = ctk.CTkToplevel(self)
        dlg.title(self.t()["form_title_student"])
        dlg.geometry("430x520")
        dlg.transient(self)
        dlg.grid_columnconfigure(0, weight=1)
        fields = ["username", "password", "student_number", "full_name", "group_id",
                  "specialty_id", "course_id", "email", "phone", "birth_date"]
        references = {
            "group_id": ("groups", False),
            "specialty_id": ("specialties", True),
            "course_id": ("courses", True),
        }
        entries: dict[str, Any] = {}
        for idx, key in enumerate(fields):
            if key in references:
                table, optional = references[key]
                self._add_reference_field(dlg, key, idx, entries, table, optional)
                continue
            entry = ctk.CTkEntry(
                dlg,
                placeholder_text=self.t().get(key, key.replace("_", " ").title()),
                show="*" if key == "password" else "",
            )
            entry.grid(row=idx, column=0, padx=18, pady=6, sticky="ew")
            entries[key] = entry

        def save() -> None:
            payload = {key: entry.get().strip() for key, entry in entries.items()}
            required = ["username", "password", "full_name", "student_number"]
            if any(not payload[key] for key in required):
                raise ValueError(self.t()["required_fields"])
            if len(payload["password"]) < 8:
                raise ValueError(self.t()["password_min_length"])
            group_id = self._reference_id(entries["group_id"])
            specialty_id = self._reference_id(entries["specialty_id"], optional=True)
            course_id = self._reference_id(entries["course_id"], optional=True)
            StudentService.create_student(
                self.app.conn,
                {
                    "username": payload["username"],
                    "password": payload["password"],
                    "student_number": payload["student_number"],
                    "full_name": payload["full_name"],
                    "group_id": group_id,
                    "specialty_id": specialty_id,
                    "course_id": course_id,
                    "email": payload.get("email") or None,
                    "phone": payload.get("phone") or None,
                    "birth_date": payload.get("birth_date") or None,
                },
            )
            self._defer_dialog_close(dlg, lambda: self.show_view("students"))
        ctk.CTkButton(
            dlg, text=self.t()["save"], command=lambda: self._run_form(save),
            fg_color="#176B87",
        ).grid(row=len(fields), column=0, padx=18, pady=12, sticky="ew")

    def add_or_edit_grade(self) -> None:
        dlg = ctk.CTkToplevel(self)
        dlg.title(self.t()["form_title_grade"])
        dlg.geometry("420x420")
        dlg.transient(self)
        dlg.grid_columnconfigure(0, weight=1)
        fields = ["student_id", "subject_id", "semester_id", "current_score",
                  "examination_score", "final_score", "note"]
        references = {
            "student_id": ("students", False),
            "subject_id": ("subjects", False),
            "semester_id": ("semesters", False),
        }
        entries: dict[str, Any] = {}
        for idx, key in enumerate(fields):
            if key in references:
                self._add_reference_field(
                    dlg, key, idx, entries, references[key][0]
                )
                continue
            label = self.t().get(key, key.replace("_", " ").title())
            if key == "final_score":
                label += " (optional)"
            entry = ctk.CTkEntry(dlg, placeholder_text=label)
            entry.grid(row=idx, column=0, padx=18, pady=6, sticky="ew")
            entries[key] = entry

        def save() -> None:
            payload = {key: entry.get().strip() for key, entry in entries.items()}
            if any(not payload[key] for key in fields[3:-2]):
                raise ValueError(self.t()["required_fields"])
            payload["student_id"] = self._reference_id(entries["student_id"])
            payload["subject_id"] = self._reference_id(entries["subject_id"])
            payload["semester_id"] = self._reference_id(entries["semester_id"])
            payload["current_score"] = float(payload["current_score"])
            payload["examination_score"] = float(payload["examination_score"])
            if payload.get("final_score"):
                payload["final_score"] = float(payload["final_score"])
            else:
                payload["final_score"] = None
            GradeService.upsert(self.app.conn, payload)
            self._defer_dialog_close(dlg, lambda: self.show_view("grades"))
        ctk.CTkButton(
            dlg, text=self.t()["save"], command=lambda: self._run_form(save),
            fg_color="#176B87",
        ).grid(row=len(fields), column=0, padx=18, pady=12, sticky="ew")

    def add_or_edit_schedule(self) -> None:
        dlg = ctk.CTkToplevel(self)
        dlg.title(self.t()["form_title_schedule"])
        dlg.geometry("420x460")
        dlg.transient(self)
        dlg.grid_columnconfigure(0, weight=1)
        fields = ["group_id", "subject_id", "weekday", "starts_at", "ends_at",
                  "teacher", "classroom", "semester_id"]
        references = {
            "group_id": ("groups", False),
            "subject_id": ("subjects", False),
            "semester_id": ("semesters", True),
        }
        entries: dict[str, Any] = {}
        for idx, key in enumerate(fields):
            if key in references:
                table, optional = references[key]
                self._add_reference_field(dlg, key, idx, entries, table, optional)
                continue
            entry = ctk.CTkEntry(
                dlg, placeholder_text=self.t().get(key, key.replace("_", " ").title())
            )
            entry.grid(row=idx, column=0, padx=18, pady=6, sticky="ew")
            entries[key] = entry

        def save() -> None:
            payload = {key: entry.get().strip() for key, entry in entries.items()}
            required_fields = ["weekday", "starts_at", "ends_at", "teacher", "classroom"]
            if any(not payload[key] for key in required_fields):
                raise ValueError(self.t()["required_fields"])
            datetime.strptime(payload["starts_at"], "%H:%M")
            datetime.strptime(payload["ends_at"], "%H:%M")
            if int(payload["weekday"]) not in range(7):
                raise ValueError(self.t()["invalid_value"])
            group_id = self._reference_id(entries["group_id"])
            subject_id = self._reference_id(entries["subject_id"])
            semester_id = self._reference_id(entries["semester_id"], optional=True)
            self.app.conn.execute(
                "INSERT INTO schedule(group_id, subject_id, weekday, starts_at, ends_at, teacher, classroom, semester_id) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    group_id,
                    subject_id,
                    int(payload["weekday"]),
                    payload["starts_at"],
                    payload["ends_at"],
                    payload["teacher"],
                    payload["classroom"],
                    semester_id,
                ),
            )
            self.app.conn.commit()
            self._defer_dialog_close(dlg, lambda: self.show_view("schedule"))
        ctk.CTkButton(
            dlg, text=self.t()["save"], command=lambda: self._run_form(save),
            fg_color="#176B87",
        ).grid(row=len(fields), column=0, padx=18, pady=12, sticky="ew")

    def add_or_edit_attendance(self) -> None:
        dlg = ctk.CTkToplevel(self)
        dlg.title(self.t()["form_title_attendance"])
        dlg.geometry("420x360")
        dlg.transient(self)
        dlg.grid_columnconfigure(0, weight=1)
        fields = ["student_id", "subject_id", "lesson_date", "status", "notes"]
        entries: dict[str, Any] = {}
        status_values = {
            self.t()["attendance_present"]: "present",
            self.t()["attendance_absent"]: "absent",
            self.t()["attendance_excused"]: "excused",
        }
        for idx, key in enumerate(fields):
            if key in {"student_id", "subject_id"}:
                self._add_reference_field(
                    dlg, key, idx, entries, "students" if key == "student_id" else "subjects"
                )
                continue
            if key == "status":
                selector = ctk.CTkOptionMenu(
                    dlg,
                    values=list(status_values),
                    fg_color="#176B87",
                    button_color="#12566D",
                    height=40,
                )
                selector.set(self.t()["attendance_present"])
                selector.grid(row=idx, column=0, padx=18, pady=6, sticky="ew")
                entries[key] = selector
                continue
            label = self.t().get(key, self.t()["note"])
            entry = ctk.CTkEntry(
                dlg,
                placeholder_text=label,
                placeholder_text_color="#77869A",
            )
            if key == "lesson_date":
                entry.insert(0, date.today().isoformat())
            entry.grid(row=idx, column=0, padx=18, pady=6, sticky="ew")
            entries[key] = entry

        def save() -> None:
            payload = {key: entry.get().strip() for key, entry in entries.items()}
            if any(not payload[key] for key in ("lesson_date",)):
                raise ValueError(self.t()["required_fields"])
            payload["student_id"] = self._reference_id(entries["student_id"])
            payload["subject_id"] = self._reference_id(entries["subject_id"])
            payload["status"] = status_values[payload["status"]]
            date.fromisoformat(payload["lesson_date"])
            self.app.conn.execute(
                "INSERT INTO attendance(student_id, subject_id, lesson_date, status, notes) VALUES (?, ?, ?, ?, ?)",
                (
                    int(payload["student_id"]),
                    int(payload["subject_id"]),
                    payload["lesson_date"],
                    payload["status"],
                    payload.get("notes"),
                ),
            )
            self.app.conn.commit()
            self._defer_dialog_close(dlg, lambda: self.show_view("attendance"))
        ctk.CTkButton(
            dlg, text=self.t()["save"], command=lambda: self._run_form(save),
            fg_color="#176B87",
        ).grid(row=len(fields), column=0, padx=18, pady=12, sticky="ew")

    def add_or_edit_payment(self) -> None:
        dlg = ctk.CTkToplevel(self)
        dlg.title(self.t()["form_title_payment"])
        dlg.geometry("420x360")
        dlg.transient(self)
        dlg.grid_columnconfigure(0, weight=1)
        fields = ["contract_id", "amount", "payment_date", "receipt_number", "note"]
        entries: dict[str, Any] = {}
        for idx, key in enumerate(fields):
            if key == "contract_id":
                self._add_reference_field(dlg, key, idx, entries, "contracts")
                continue
            entry = ctk.CTkEntry(
                dlg, placeholder_text=self.t().get(key, self.t()["note"])
            )
            if key == "payment_date":
                entry.insert(0, date.today().isoformat())
            entry.grid(row=idx, column=0, padx=18, pady=6, sticky="ew")
            entries[key] = entry

        def save() -> None:
            payload = {key: entry.get().strip() for key, entry in entries.items()}
            if any(not payload[key] for key in fields[1:-1]):
                raise ValueError(self.t()["required_fields"])
            payload["contract_id"] = self._reference_id(entries["contract_id"])
            date.fromisoformat(payload["payment_date"])
            PaymentService.create_payment(self.app.conn, {
                "contract_id": int(payload["contract_id"]),
                "amount": float(payload["amount"]),
                "payment_date": payload["payment_date"],
                "receipt_number": payload["receipt_number"],
                "note": payload.get("note") or None,
            })
            self._defer_dialog_close(dlg, lambda: self.show_view("payments"))
        ctk.CTkButton(
            dlg, text=self.t()["save"], command=lambda: self._run_form(save),
            fg_color="#176B87",
        ).grid(row=len(fields), column=0, padx=18, pady=12, sticky="ew")

    def add_or_edit_message(self) -> None:
        dlg = ctk.CTkToplevel(self)
        dlg.title(self.t()["form_title_message"])
        dlg.geometry("420x360")
        dlg.transient(self)
        dlg.grid_columnconfigure(0, weight=1)
        fields = ["target_group_id", "title", "body"]
        entries: dict[str, Any] = {}
        for idx, key in enumerate(fields):
            if key == "target_group_id":
                self._add_reference_field(
                    dlg, key, idx, entries, "groups", optional=True
                )
                continue
            entry = ctk.CTkEntry(
                dlg, placeholder_text=self.t().get(key, key.replace("_", " ").title())
            )
            entry.grid(row=idx, column=0, padx=18, pady=6, sticky="ew")
            entries[key] = entry

        def save() -> None:
            payload = {key: entry.get().strip() for key, entry in entries.items()}
            if not payload["title"] or not payload["body"]:
                raise ValueError(self.t()["required_fields"])
            payload["target_group_id"] = self._reference_id(
                entries["target_group_id"], optional=True
            )
            self.app.conn.execute(
                "INSERT INTO messages(author_id, target_group_id, title, body, published_at, is_published) VALUES (?, ?, ?, ?, CURRENT_TIMESTAMP, 1)",
                (self.user["id"], payload["target_group_id"], payload["title"], payload["body"]),
            )
            self.app.conn.commit()
            self._defer_dialog_close(dlg, lambda: self.show_view("messages"))
        ctk.CTkButton(
            dlg, text=self.t()["save"], command=lambda: self._run_form(save),
            fg_color="#176B87",
        ).grid(row=len(fields), column=0, padx=18, pady=12, sticky="ew")


class StudentCabinetApp(ctk.CTk):
    """Application root that swaps login and dashboard by role."""

    def __init__(self) -> None:
        super().__init__()
        ctk.set_appearance_mode("light")
        ctk.set_default_color_theme("blue")
        self.title(APP_TITLE)
        self.geometry("1100x700")
        self.minsize(900, 620)
        self.db = DatabaseManager()
        self.conn = self.db.conn
        self.current_user: dict[str, Any] | None = None
        self.language = CURRENT_LANG
        self.dashboard: ctk.CTkFrame | None = None
        self.login_window: LoginWindow | None = None
        self.protocol("WM_DELETE_WINDOW", self.on_login_close)
        self.show_login()

    def safe_destroy(self, widget: Any | None) -> None:
        if widget is not None and hasattr(widget, "winfo_exists") and widget.winfo_exists():
            widget.destroy()

    def show_login(self) -> None:
        if self.dashboard is not None:
            self.safe_destroy(self.dashboard)
            self.dashboard = None
        if self.login_window is not None:
            self.safe_destroy(self.login_window)
            self.login_window = None
        self.geometry("1100x700")
        self.minsize(900, 620)
        self.login_window = LoginWindow(self, self.login_callback, self.language)
        self.login_window.pack(fill="both", expand=True)

    def on_login_close(self) -> None:
        if self.login_window is not None:
            self.safe_destroy(self.login_window)
            self.login_window = None
        self.db.close()
        self.destroy()

    def login_callback(self, username: str, password: str, language: str) -> None:
        user = self.conn.execute("SELECT * FROM users WHERE username = ? AND is_active = 1", (username,)).fetchone()
        if user is None or not AuthenticationService.verify_password(user["password_hash"], password):
            raise ValueError("Invalid username or password")
        self.language = language
        self.current_user = {"id": user["id"], "username": user["username"], "role": user["role"]}
        if self.login_window is not None:
            self.safe_destroy(self.login_window)
            self.login_window = None
        self.geometry("1180x760")
        self.minsize(980, 640)
        self.title(TRANSLATIONS[language]["app_name"])
        dashboard_class = AdminDashboard if user["role"] == "admin" else StudentDashboard
        self.dashboard = dashboard_class(self, self, self.current_user, language)

    def logout(self) -> None:
        self.current_user = None
        if self.dashboard is not None:
            self.safe_destroy(self.dashboard)
            self.dashboard = None
        self.show_login()


def run_client() -> None:
    app = StudentCabinetApp()
    app.mainloop()


def main() -> None:
    parser = argparse.ArgumentParser(description="Student Cabinet")
    parser.add_argument("mode", nargs="?", choices=["client", "server"], default="client", help="Run the desktop app. Server mode is intentionally unavailable in this local-only build.")
    args = parser.parse_args()
    if args.mode == "server":
        print("This build is local-only and uses SQLite. Run: python main.py client")
        return
    run_client()


if __name__ == "__main__":
    main()
